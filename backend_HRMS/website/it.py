from datetime import datetime
import uuid

from flask import Blueprint, current_app, jsonify, request, send_file
from sqlalchemy import extract, func
from sqlalchemy.orm import defer, load_only, selectinload
from .datetime_utils import utc_now, isoformat_api
from flask_jwt_extended import get_jwt_identity, jwt_required, get_jwt

from . import db
from .models.Admin_models import Admin
from .query_utils import (
    FULL_LIST,
    apply_search,
    ci_equals,
    empty_pagination,
    list_params,
    list_rows,
    pagination_enabled,
    rows_sorted_by_id,
    sort_order,
)
from .tenant_context import (
    current_admin,
    get_admin_in_tenant,
    get_current_tenant_id,
    tenant_admin_ids,
    tenant_admin_query,
)
from .itam.tenancy import it_get, it_query, it_tenant_clause
from .itam.media import persist_itam_media, present_itam_media
from .models.it_models import (
    ITAssetAssignment,
    ITAssetReturnRequest,
    ITAssetUnit,
    ITDeviceType,
    ITDeletedAssetLog,
    ITInventoryItem,
    ITInventoryQuantityAssignment,
    ITOfficeStockDeployment,
    ITParcelExport,
    ITParcelExportItem,
    ITParcelImport,
    ITRepairBill,
    ITRemovedAsset,
    ITSoftwareLicense,
    ITSupportTicket,
)
from .email import (
    send_it_assignment_notification,
    send_it_return_request_email,
    send_it_return_request_status_email,
)
from .noc_department_service import download_noc_document, list_noc_requests, upload_noc_document
from .employee_photo import photo_url_for_admin, prime_photo_urls


it_bp = Blueprint("it", __name__)


def _stored_media(values, folder):
    """(keys, error response). Error is None when the photos could be stored."""
    try:
        return persist_itam_media(values, folder), None
    except ValueError as exc:
        return None, _err(str(exc))


def _shown_media(values):
    return present_itam_media(values)


def _allow_self_asset_access_for_non_it():
    """My Assets: employees read own records and submit return requests (ownership checked in handlers)."""
    from .plan_features import has_feature

    if not has_feature("dashboard_my_assets"):
        return False

    actor = _current_admin()
    if not actor:
        return False

    endpoint = request.endpoint or ""

    if request.method == "GET" and endpoint in {"it.employee_assets", "it.employee_return_requests"}:
        req_emp = (request.view_args or {}).get("emp_id")
        if not req_emp:
            return False
        target = _resolve_admin_by_emp_key(req_emp)
        if not target:
            return False
        return int(actor.id) == int(target.id)

    if request.method == "POST" and endpoint == "it.create_return_request":
        return True

    return False


def _allow_finance_purchase_access(claims):
    """Finance reviews purchase budget on IT-mounted purchase routes without an IT seat."""
    from .plan_features import can_access_finance_panel

    if not can_access_finance_panel(claims):
        return False
    endpoint = request.endpoint or ""
    return endpoint in {
        "it.inventory_purchase_list",
        "it.inventory_purchase_finance",
    } and request.method in ("GET", "POST", "OPTIONS")


def _allow_manager_purchase_access():
    """Mapped managers act on purchase requests without an IT seat."""
    endpoint = request.endpoint or ""
    return endpoint in {
        "it.inventory_purchase_list",
        "it.inventory_purchase_manager",
    } and request.method in ("GET", "POST", "OPTIONS")


@it_bp.before_request
def _it_plan_guard():
    from flask_jwt_extended import get_jwt, verify_jwt_in_request
    from .plan_features import can_access_it_panel, plan_forbidden_response

    if request.method == "OPTIONS":
        return None

    try:
        verify_jwt_in_request()
    except Exception:
        return jsonify({"success": False, "message": "Unauthorized"}), 401

    claims = get_jwt() or {}
    if not can_access_it_panel(claims):
        if _allow_self_asset_access_for_non_it():
            return None
        if _allow_finance_purchase_access(claims):
            return None
        if _allow_manager_purchase_access():
            return None
        return plan_forbidden_response("it_panel")
    return None


def _ok(payload=None, message="OK", code=200):
    body = {"success": True, "message": message}
    if payload:
        body.update(payload)
    return jsonify(body), code


def _err(message="Bad request", code=400):
    return jsonify({"success": False, "message": message}), code


def _list_ok(key, items, pagination):
    """Old list body; ``pagination`` is added only when the caller asked for a page."""
    payload = {key: items}
    if pagination is not None:
        payload["pagination"] = pagination
    return _ok(payload)


def _sort_or_err(params, allowed, model, default):
    """(order_by, error response) for the request's ``sort``."""
    order = sort_order(params.sort, allowed, model.id, default)
    if order is None:
        names = ", ".join(sorted(allowed))
        return None, _err(f"sort must be one of: {names} (prefix with - for descending)")
    return order, None


def _record_transition_or_err(**kwargs):
    """
    Append an activity-log transition. Remarks are enforced only when
    itam_transitions_v1 is ON. Returns Flask error tuple on validation failure.
    """
    from flask import current_app
    from .itam.transition_service import TransitionValidationError, record_transition

    try:
        record_transition(config=current_app.config, **kwargs)
        return None
    except TransitionValidationError as exc:
        return _err(str(exc), 400)


def _apply_lifecycle_or_err(
    unit,
    *,
    lifecycle_status,
    custody_type=None,
    custody=None,
    actor_admin_id=None,
    dual_write_legacy=True,
):
    """Apply P3 lifecycle dual-write when itam_lifecycle_v1 is ON."""
    from flask import current_app
    from .itam.lifecycle_service import (
        LifecycleValidationError,
        apply_unit_lifecycle,
        lifecycle_enabled,
    )

    if not lifecycle_enabled(current_app.config):
        return None
    try:
        apply_unit_lifecycle(
            unit,
            lifecycle_status=lifecycle_status,
            custody_type=custody_type,
            custody=custody,
            actor_admin_id=actor_admin_id,
            dual_write_legacy=dual_write_legacy,
            config=current_app.config,
        )
        return None
    except LifecycleValidationError as exc:
        return _err(str(exc), 400)


def _iso(dt):
    return isoformat_api(dt)


def _parse_dt(value):
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except Exception:
        return None


def _parse_date(value):
    if not value:
        return None
    try:
        return datetime.strptime(str(value), "%Y-%m-%d").date()
    except Exception:
        return None


def _ensure_parcel_name_columns_runtime():
    """Best-effort runtime safety: add free-text parcel columns if missing."""
    try:
        from sqlalchemy import inspect, text

        insp = inspect(db.engine)
        tables = set(insp.get_table_names())
        dialect = db.engine.dialect.name

        def addcol(table, col):
            if table not in tables:
                return
            existing = {c["name"] for c in insp.get_columns(table)}
            if col in existing:
                return
            if dialect == "postgresql":
                stmt = text(f'ALTER TABLE "{table}" ADD COLUMN {col} VARCHAR(120) NULL')
            else:
                stmt = text(f"ALTER TABLE {table} ADD COLUMN {col} VARCHAR(120) NULL")
            with db.engine.begin() as conn:
                conn.execute(stmt)
            current_app.logger.info("Added missing parcel column %s.%s at runtime", table, col)

        addcol("it_parcel_imports", "received_by_name")
        addcol("it_parcel_exports", "exported_by_name")
        addcol("it_parcel_exports", "inventory_category")
        addcol("it_parcel_export_items", "make")
        addcol("it_asset_units", "project_code")
        addcol("it_asset_units", "device_location")
        addcol("it_asset_units", "device_type")
    except Exception as e:
        current_app.logger.warning("parcel runtime column ensure skipped: %s", e)


def _current_admin():
    return current_admin()


def _admin_name(admin):
    if not admin:
        return None
    return (admin.first_name or "").strip() or admin.email or f"Admin-{admin.id}"


def _serialize_admin_employee(admin, assigned_assets=None):
    photo = photo_url_for_admin(admin)
    return {
        "id": admin.emp_id or str(admin.id),
        "empId": admin.emp_id or str(admin.id),
        "adminId": admin.id,
        "name": admin.first_name or "",
        "type": admin.emp_type or "",
        "circle": admin.circle or "",
        "email": admin.email or "",
        "photo": photo,
        "photo_url": photo,
        "activated": bool(admin.is_active),
        "assignedAssets": assigned_assets if assigned_assets is not None else _serialize_emp_assets(admin),
    }


def _resolve_admin_by_emp_key(emp_key):
    """Resolve Admin from URL/key: emp_id, numeric admin id, or email."""
    key = str(emp_key or "").strip()
    if not key:
        return None
    emp = tenant_admin_query().filter(
        db.func.lower(db.func.coalesce(Admin.emp_id, "")) == key.lower()
    ).first()
    if emp:
        return emp
    if key.isdigit():
        return get_admin_in_tenant(int(key))
    return tenant_admin_query().filter(
        db.func.lower(db.func.coalesce(Admin.email, "")) == key.lower()
    ).first()


def _canonical_inv_cat(value):
    raw = (value or "").strip()
    if raw.lower() in ("office assets", "fixed assets"):
        return "Fixed Assets"
    return raw


def _inv_cat_aliases(value):
    canon = _canonical_inv_cat(value)
    if canon == "Fixed Assets":
        return ("Fixed Assets", "Office Assets")
    return (canon,) if canon else ()


def _is_qty_managed_inventory(category, inventory_category=None):
    cat = (category or "").strip().lower()
    if cat in ("accessories", "consumables", "stock"):
        return True
    inv = _canonical_inv_cat(inventory_category)
    if inv in ("Fixed Assets", "Infrastructure Assets"):
        return cat == "stock" or cat in ("accessories", "consumables")
    return False


def _coerce_optional_int(value):
    """Return int id or None — rejects frontend synthetic tokens like inv-slot-12-0."""
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    try:
        s = str(value).strip()
        if not s or not s.lstrip("-").isdigit():
            return None
        return int(s)
    except (TypeError, ValueError):
        return None


def _existing_fk_id(model, value):
    """Return value only when the FK target row exists in the caller's tenant; otherwise None (avoids MySQL 1452)."""
    pk = _coerce_optional_int(value)
    if pk is None:
        return None
    found = get_admin_in_tenant(pk) if model is Admin else it_get(model, pk)
    return pk if found is not None else None


def _inventory_id_from_export_slot(value):
    """Parse inv-slot-{inventoryId}-{index} from Ready for Export UI."""
    s = str(value or "").strip()
    if not s.startswith("inv-slot-"):
        return None
    parts = s.split("-")
    if len(parts) >= 4:
        return _coerce_optional_int(parts[2])
    return None


def _resolve_parcel_export_item_ids(it):
    """Map export payload row to real asset_unit_id and optional inventory_item_id."""
    raw_unit = it.get("asset_unit_id")
    # Prefer explicit unit id only. Do not treat inventory slot tokens (inv-slot-*) as unit ids.
    if raw_unit is None:
        candidate = it.get("id")
        if _coerce_optional_int(candidate) is not None:
            raw_unit = candidate
    asset_unit_id = _existing_fk_id(ITAssetUnit, raw_unit)

    inventory_item_id = _coerce_optional_int(it.get("inventory_item_id"))
    if inventory_item_id is None:
        inventory_item_id = _inventory_id_from_export_slot(it.get("id"))
    if inventory_item_id is not None and it_get(ITInventoryItem, inventory_item_id) is None:
        inventory_item_id = None
    return asset_unit_id, inventory_item_id


def _effective_inventory_category(item):
    """Align stored inventory_category with item.category (fixes legacy mis-tagged rows)."""
    ic = (item.inventory_category or "").strip()
    cat = (item.category or "").strip().lower()
    if cat == "vehicle":
        return "Transport Assets"
    if cat == "equipment":
        return "Infrastructure Assets"
    if cat == "stock":
        if _canonical_inv_cat(ic) in ("Fixed Assets", "Infrastructure Assets", "Transport Assets"):
            return _canonical_inv_cat(ic)
        return ic
    if cat in ("hardware", "software", "accessories", "consumables"):
        return "IT Assets"
    return _canonical_inv_cat(ic) or "IT Assets"


def _serialize_inventory_item(item, summary=False):
    payload = {
        "id": item.id,
        "inventory_code": item.inventory_code,
        "name": item.name,
        "category": item.category,
        "inventory_category": _effective_inventory_category(item),
        "hw_type": item.hw_type,
        "vendor": item.vendor or "",
        "purchase_date": item.purchase_date.isoformat() if item.purchase_date else None,
        "location": item.location or "",
        "notes": item.notes or "",
        "totalQuantity": int(item.total_quantity or 0),
        "availableQuantity": int(item.available_quantity or 0),
        "assignedQuantity": int(item.assigned_quantity or 0),
        "notWorkingQuantity": int(item.not_working_quantity or 0),
        "repairQuantity": int(item.repair_quantity or 0),
        "created_at": _iso(item.created_at),
        "updated_at": _iso(item.updated_at),
    }
    if not summary:
        payload["photos"] = _shown_media(item.photos_json or [])
        payload["receipts"] = _shown_media(item.receipts_json or [])
    return payload


_INVENTORY_MEDIA = (defer(ITInventoryItem.photos_json), defer(ITInventoryItem.receipts_json))


def _serialize_asset_unit(unit, summary=False):
    from flask import current_app, has_app_context
    from .itam.lifecycle_service import lifecycle_enabled, serialize_lifecycle_fields

    assigned_admin = unit.assigned_to_admin
    payload = {
        "id": unit.id,
        "unitCode": unit.unit_code,
        "inventoryId": unit.inventory_item_id,
        "assetName": unit.asset_name,
        "category": unit.category,
        "hwType": unit.hw_type,
        "brand": unit.brand,
        "make": unit.make,
        "model": unit.model,
        "serialNumber": unit.serial_number,
        "deviceType": getattr(unit, "device_type", None),
        "imei1": unit.imei1,
        "imei2": unit.imei2,
        "projectCode": unit.project_code,
        "deviceLocation": unit.device_location,
        "status": unit.status,
        "isDailyPool": bool(getattr(unit, "is_daily_pool", False)),
        "assignedTo": unit.assigned_to_admin_id,
        "assignedToName": _admin_name(unit.assigned_to_admin),
        "assignedToEmpId": (assigned_admin.emp_id if assigned_admin else None),
        "assignedToPhoto": photo_url_for_admin(assigned_admin),
        "assignedDate": _iso(unit.assigned_at),
        "assetTag": unit.asset_tag,
        "exportedTo": unit.exported_to,
        "exportedAt": _iso(unit.exported_at),
        "repairDate": _iso(unit.repair_date),
        "created_at": _iso(unit.created_at),
        "updated_at": _iso(unit.updated_at),
    }
    if not summary:
        payload["photos"] = _shown_media(unit.photos_json or [])
        payload["assignmentPhotos"] = _shown_media(unit.assignment_photos_json or [])
    if has_app_context() and lifecycle_enabled(current_app.config):
        payload.update(serialize_lifecycle_fields(unit))
    return payload


def units_for_listing(query, order_by, params=FULL_LIST):
    """(units, pagination) for _serialize_asset_unit lists: id-only sort, related rows loaded in batches."""
    from .itam.lifecycle_service import lifecycle_enabled, prime_open_deployments

    lifecycle_on = lifecycle_enabled(current_app.config)
    options = [selectinload(ITAssetUnit.assigned_to_admin)]
    if lifecycle_on:
        options.append(selectinload(ITAssetUnit.inventory_item))
    if params.summary:
        options += [defer(ITAssetUnit.photos_json), defer(ITAssetUnit.assignment_photos_json)]
    rows, pagination = list_rows(query, ITAssetUnit, order_by, params, options)
    prime_photo_urls(r.assigned_to_admin_id for r in rows)
    if lifecycle_on:
        prime_open_deployments(rows)
    return rows, pagination


def load_units_for_listing(query, order_by):
    return units_for_listing(query, order_by)[0]


_UNIT_SEARCH = (
    ITAssetUnit.asset_name,
    ITAssetUnit.unit_code,
    ITAssetUnit.serial_number,
    ITAssetUnit.device_type,
    ITAssetUnit.brand,
    ITAssetUnit.make,
    ITAssetUnit.model,
    ITAssetUnit.hw_type,
    ITAssetUnit.project_code,
    ITAssetUnit.imei1,
    ITAssetUnit.imei2,
    ITAssetUnit.asset_tag,
    ITAssetUnit.device_location,
)


def _serialize_license(lic):
    assigned_admin = lic.assigned_to_admin
    return {
        "id": lic.id,
        "licenseCode": lic.license_code,
        "inventoryId": lic.inventory_item_id,
        "name": lic.name,
        "subscriptionStart": lic.subscription_start.isoformat() if lic.subscription_start else None,
        "subscriptionEnd": lic.subscription_end.isoformat() if lic.subscription_end else None,
        "status": lic.status,
        "assignedTo": lic.assigned_to_admin_id,
        "assignedToName": _admin_name(lic.assigned_to_admin),
        "assignedToEmpId": (assigned_admin.emp_id if assigned_admin else None),
        "assignedToPhoto": photo_url_for_admin(assigned_admin),
        "assignedDate": _iso(lic.assigned_at),
        "created_at": _iso(lic.created_at),
        "updated_at": _iso(lic.updated_at),
    }


def _serialize_ticket(t):
    return {
        "id": t.id,
        "ticketCode": t.ticket_code,
        "requesterAdminId": t.requester_admin_id,
        "requesterName": _admin_name(t.requester_admin),
        "assigneeAdminId": t.assignee_admin_id,
        "assigneeName": _admin_name(t.assignee_admin),
        "title": t.title,
        "description": t.description,
        "priority": t.priority,
        "status": t.status,
        "resolvedAt": _iso(t.resolved_at),
        "created_at": _iso(t.created_at),
        "updated_at": _iso(t.updated_at),
    }


def _serialize_removed_asset(r, summary=False):
    payload = {
        "id": r.id,
        "removedCode": r.removed_code,
        "assetUnitId": r.asset_unit_id,
        "inventoryId": r.inventory_item_id,
        "ownerAdminId": r.owner_admin_id,
        "ownerName": _admin_name(r.owner_admin),
        "removedByAdminId": r.removed_by_admin_id,
        "removedByName": _admin_name(r.removed_by_admin),
        "name": r.name,
        "category": r.category,
        "reason": r.reason,
        "removedAt": _iso(r.removed_at),
    }
    if not summary:
        payload["photos"] = _shown_media(r.photos_json or [])
    return payload


def _serialize_deleted_log(d):
    return {
        "id": d.id,
        "deleteCode": d.delete_code,
        "assetUnitId": d.asset_unit_id,
        "inventoryId": d.inventory_item_id,
        "deletedByAdminId": d.deleted_by_admin_id,
        "deletedByName": (d.deleted_by_name or "").strip() or _admin_name(d.deleted_by_admin),
        "assetName": d.asset_name,
        "category": d.category,
        "serialNumber": d.serial_number,
        "reason": d.reason,
        "deletedAt": _iso(d.deleted_at),
    }


def _serialize_parcel_import(i):
    return {
        "id": i.id,
        "importCode": i.import_code,
        "from": i.source,
        "assetName": i.asset_name,
        "count": i.count,
        "idNo": i.id_no,
        "receivedByAdminId": i.received_by_admin_id,
        "receivedBy": (i.received_by_name or "").strip() or _admin_name(i.received_by_admin),
        "date": _iso(i.received_at),
        "photos": _shown_media(i.photos_json or []),
    }


def _serialize_parcel_export(e):
    return {
        "id": e.id,
        "exportCode": e.export_code,
        "to": e.destination,
        "idNo": e.id_no,
        "exportedByAdminId": e.exported_by_admin_id,
        "exportedBy": (e.exported_by_name or "").strip() or _admin_name(e.exported_by_admin),
        "date": _iso(e.exported_at),
        "photos": _shown_media(e.parcel_photos_json or []),
        "inventoryCategory": (e.inventory_category or "").strip() or None,
        "count": len(e.items),
        "assets": [
            {
                "id": it.id,
                "assetUnitId": it.asset_unit_id,
                "assetName": it.asset_name,
                "serialNo": it.serial_number,
                "brand": it.brand,
                "make": (it.make or "").strip()
                or ((it.asset_unit.make or "").strip() if it.asset_unit else ""),
                "model": it.model,
                "individualPhoto": _shown_media(
                    (it.individual_photo_json or [None])[0]
                    if isinstance(it.individual_photo_json, list)
                    else it.individual_photo_json
                ),
            }
            for it in e.items
        ],
    }


def _serialize_return_request(r, summary=False):
    payload = {
        "id": r.id,
        "requestCode": r.request_code,
        "requesterAdminId": r.requester_admin_id,
        "requesterEmpId": r.requester_emp_id,
        "requesterName": _admin_name(r.requester_admin),
        "requesterEmail": (r.requester_admin.email if r.requester_admin else None),
        "assetUnitId": r.asset_unit_id,
        "softwareLicenseId": r.software_license_id,
        "inventoryItemId": r.inventory_item_id,
        "quantity": int(r.quantity or 1),
        "assetName": r.asset_name,
        "category": r.category,
        "reason": r.reason,
        "returnDestination": r.return_destination or "available",
        "status": r.status,
        "approvedByAdminId": r.approved_by_admin_id,
        "approvedByName": _admin_name(r.approved_by_admin),
        "approvedAt": _iso(r.approved_at),
        "receiptConfirmedByAdminId": r.receipt_confirmed_by_admin_id,
        "receiptConfirmedByName": _admin_name(r.receipt_confirmed_by_admin),
        "receiptConfirmedAt": _iso(r.receipt_confirmed_at),
        "rejectionReason": r.rejection_reason,
        "createdAt": _iso(r.created_at),
        "updatedAt": _iso(r.updated_at),
    }
    if not summary:
        payload["photos"] = _shown_media(r.photos_json or [])
    return payload


_RETURN_REQUEST_ADMINS = (
    selectinload(ITAssetReturnRequest.requester_admin),
    selectinload(ITAssetReturnRequest.approved_by_admin),
    selectinload(ITAssetReturnRequest.receipt_confirmed_by_admin),
)
_RETURN_REQUESTS_NEWEST = (ITAssetReturnRequest.created_at.desc(), ITAssetReturnRequest.id.desc())


def _return_requests_newest_first(query):
    return rows_sorted_by_id(query, ITAssetReturnRequest, _RETURN_REQUESTS_NEWEST, _RETURN_REQUEST_ADMINS)


def _next_code(prefix, model, field_name):
    """Next code inside the caller's company (INV-0001 starts again for each company)."""
    code_col = getattr(model, field_name)
    latest = (
        db.session.query(model.id, code_col)
        .filter(it_tenant_clause(model))
        .order_by(model.id.desc())
        .first()
    )
    if not latest:
        return f"{prefix}-0001"
    latest_id, current = latest[0], (latest[1] or "")
    try:
        n = int(str(current).split("-")[-1]) + 1
    except Exception:
        n = int(latest_id) + 1
    return f"{prefix}-{str(n).zfill(4)}"


def _paginate_parcel_rows(model, *, page, per_page):
    """
    Page by primary key id only, then load full rows.

    Parcel tables store base64 photos in JSON columns; ORDER BY on the full
    row triggers MySQL "Out of sort memory" (errno 1038). Sorting bare ids
    avoids that while preserving newest-first order.
    """
    in_tenant = it_tenant_clause(model)
    total = db.session.query(db.func.count(model.id)).filter(in_tenant).scalar() or 0
    pages = max(1, (int(total) + per_page - 1) // per_page) if total else 1
    page = min(max(1, page), pages)
    id_rows = (
        db.session.query(model.id)
        .filter(in_tenant)
        .order_by(model.id.desc())
        .offset((page - 1) * per_page)
        .limit(per_page)
        .all()
    )
    ids = [int(r[0]) for r in id_rows]
    if not ids:
        items = []
    else:
        found = model.query.filter(model.id.in_(ids)).all()
        by_id = {row.id: row for row in found}
        items = [by_id[i] for i in ids if i in by_id]
    return {
        "items": items,
        "page": page,
        "per_page": per_page,
        "total": int(total),
        "pages": pages,
        "has_next": page < pages,
        "has_prev": page > 1,
    }


def _emp_unit_asset(u):
    return {
        "id": u.id,
        "assetId": u.unit_code,
        "assetTag": u.asset_tag,
        "name": u.asset_name,
        "category": u.category or "Hardware",
        "status": "Assigned" if u.status == "assigned" else (u.status or "Assigned"),
        "hwType": u.hw_type,
        "brand": u.brand or "",
        "make": u.make or "",
        "model": u.model or "",
        "serialNumber": u.serial_number or "",
        "imei1": u.imei1,
        "imei2": u.imei2,
        "photos": _shown_media(u.assignment_photos_json or u.photos_json or []),
        "assignedDate": _iso(u.assigned_at),
    }


def _emp_license_asset(s):
    return {
        "id": s.id,
        "licenseId": s.id,
        "name": s.name,
        "category": "Software",
        "status": "Assigned" if s.status == "assigned" else (s.status or "Assigned"),
        "subscriptionStart": s.subscription_start.isoformat() if s.subscription_start else None,
        "subscriptionEnd": s.subscription_end.isoformat() if s.subscription_end else None,
        "photos": [],
        "assignedDate": _iso(s.assigned_at),
    }


def _emp_quantity_asset(qa):
    item = qa.inventory_item
    category = (item.category if item else "") or ""
    cat_norm = category.strip().lower()
    if cat_norm.startswith("consumable"):
        category = "Consumables"
    elif cat_norm.startswith("accessor"):
        category = "Accessories"
    return {
        "id": f"invq-{qa.id}",
        "inventoryAssignmentId": qa.id,
        "inventoryId": qa.inventory_item_id,
        "name": item.name if item else "Inventory item",
        "category": category or "Accessories",
        "quantity": int(qa.quantity or 0),
        "status": "Assigned",
        "photos": [],
        "assignedDate": _iso(qa.updated_at or qa.created_at),
    }


def _emp_assets_by_admin(admin_ids):
    """
    {admin_id: assigned assets} for many employees: units, then licenses, then stock,
    each in id order. ``admin_ids`` is a list of ids or a SELECT of ids.
    """
    units = rows_sorted_by_id(
        ITAssetUnit.query.filter(ITAssetUnit.assigned_to_admin_id.in_(admin_ids)),
        ITAssetUnit,
        (ITAssetUnit.id.asc(),),
    )
    licenses = (
        ITSoftwareLicense.query.filter(ITSoftwareLicense.assigned_to_admin_id.in_(admin_ids))
        .order_by(ITSoftwareLicense.id.asc())
        .all()
    )
    qty_rows = (
        ITInventoryQuantityAssignment.query.options(selectinload(ITInventoryQuantityAssignment.inventory_item))
        .filter(
            ITInventoryQuantityAssignment.assigned_to_admin_id.in_(admin_ids),
            ITInventoryQuantityAssignment.quantity > 0,
        )
        .order_by(ITInventoryQuantityAssignment.id.asc())
        .all()
    )
    out = {}
    for u in units:
        out.setdefault(u.assigned_to_admin_id, []).append(_emp_unit_asset(u))
    for s in licenses:
        out.setdefault(s.assigned_to_admin_id, []).append(_emp_license_asset(s))
    for qa in qty_rows:
        out.setdefault(qa.assigned_to_admin_id, []).append(_emp_quantity_asset(qa))
    return out


def _serialize_emp_assets(emp):
    return _emp_assets_by_admin([emp.id]).get(emp.id, [])


def _serialize_admin_employees(admins):
    """_serialize_admin_employee for many admins with batched asset and photo lookups."""
    assets = _emp_assets_by_admin([a.id for a in admins]) if admins else {}
    prime_photo_urls(a.id for a in admins)
    return [_serialize_admin_employee(a, assets.get(a.id, [])) for a in admins]


def _holds_any_asset():
    """Admin holds a unit, a license seat or a non-zero stock quantity."""
    return db.or_(
        Admin.id.in_(
            db.select(ITAssetUnit.assigned_to_admin_id).where(ITAssetUnit.assigned_to_admin_id.isnot(None))
        ),
        Admin.id.in_(
            db.select(ITSoftwareLicense.assigned_to_admin_id).where(
                ITSoftwareLicense.assigned_to_admin_id.isnot(None)
            )
        ),
        Admin.id.in_(
            db.select(ITInventoryQuantityAssignment.assigned_to_admin_id).where(
                ITInventoryQuantityAssignment.quantity > 0
            )
        ),
    )


_EMPLOYEE_SEARCH = (Admin.first_name, Admin.email, Admin.emp_id)
_EMPLOYEE_SORTS = {"name": Admin.first_name, "emp_id": Admin.emp_id, "email": Admin.email}


@it_bp.route("/employees/assigned-assets", methods=["GET"])
@jwt_required()
def list_employee_assigned_assets():
    params = list_params()
    if params.page is not None or params.search or params.sort:
        order, err = _sort_or_err(params, _EMPLOYEE_SORTS, Admin, (Admin.id.asc(),))
        if err:
            return err
        q = apply_search(tenant_admin_query().filter(_holds_any_asset()), _EMPLOYEE_SEARCH, params.search)
        rows, pagination = list_rows(q, Admin, order, params, sort_ids_first=False)
        return _list_ok("employees", _serialize_admin_employees(rows), pagination)

    assets = _emp_assets_by_admin(tenant_admin_ids())
    if not assets:
        return _ok({"employees": []})
    rows = tenant_admin_query().order_by(Admin.id.asc()).all()
    prime_photo_urls(emp.id for emp in rows if assets.get(emp.id))
    out = []
    for emp in rows:
        emp_assets = assets.get(emp.id)
        if not emp_assets:
            continue
        out.append(_serialize_admin_employee(emp, emp_assets))
    return _ok({"employees": out})


def _recalc_inventory_counts(inventory_id):
    item = it_get(ITInventoryItem, inventory_id)
    if not item:
        return
    if item.category.lower() == "software":
        seats = [
            s
            for (s,) in db.session.query(ITSoftwareLicense.status).filter(
                ITSoftwareLicense.inventory_item_id == inventory_id
            )
        ]
        item.total_quantity = len(seats)
        item.available_quantity = len([s for s in seats if s == "available"])
        item.assigned_quantity = len([s for s in seats if s == "assigned"])
        item.not_working_quantity = 0
        item.repair_quantity = 0
    else:
        units = [
            u
            for (u,) in db.session.query(ITAssetUnit.status).filter(
                ITAssetUnit.inventory_item_id == inventory_id
            )
        ]
        item.total_quantity = len(units)
        item.available_quantity = len([u for u in units if (u or "").lower() == "available"])
        item.assigned_quantity = len(
            [u for u in units if (u or "").lower() in ("assigned", "daily_out")]
        )
        item.not_working_quantity = len([u for u in units if u in ("not-working", "notWorking")])
        item.repair_quantity = len([u for u in units if u == "repair"])


@it_bp.route("/summary", methods=["GET"])
@jwt_required()
def it_summary():
    total_assets = it_query(ITAssetUnit).count()
    open_tickets = it_query(ITSupportTicket).filter(ITSupportTicket.status == "pending").count()
    assigned_assets = it_query(ITAssetUnit).filter(ITAssetUnit.status == "assigned").count()
    day_use_assets = it_query(ITAssetUnit).filter(
        db.func.lower(ITAssetUnit.status) == "daily_out"
    ).count()
    available_assets = it_query(ITAssetUnit).filter(ITAssetUnit.status == "available").count()
    available_software = it_query(ITSoftwareLicense).filter(ITSoftwareLicense.status == "available").count()
    pending_returns = it_query(ITAssetReturnRequest).filter(
        db.func.lower(ITAssetReturnRequest.status) == "pending"
    ).count()

    pending_emails = 0
    pending_assignments = 0
    pending_noc = 0
    pending_inventory_noc = 0
    pending_day_use = 0
    try:
        from .models.daily_checkout import DailyCheckoutRequest
        from .daily_checkout.state import OPEN_STATUSES

        pending_day_use = it_query(DailyCheckoutRequest).filter(
            DailyCheckoutRequest.status.in_(tuple(OPEN_STATUSES))
        ).count()
    except Exception:
        pending_day_use = 0
    try:
        from .models.official_email_request import OfficialEmailRequest
        from .tenant_context import get_current_tenant_id

        tid = get_current_tenant_id()
        pending_emails = OfficialEmailRequest.query.filter(
            OfficialEmailRequest.tenant_id == tid,
            db.func.lower(OfficialEmailRequest.status) == "pending",
        ).count()
    except Exception:
        pending_emails = 0
    try:
        from .models.asset_assignment_request import AssetAssignmentRequest
        from .tenant_context import get_current_tenant_id

        tid = get_current_tenant_id()
        pending_assignments = AssetAssignmentRequest.query.filter(
            AssetAssignmentRequest.tenant_id == tid,
            db.func.lower(AssetAssignmentRequest.status) == "pending",
        ).count()
    except Exception:
        pending_assignments = 0
    try:
        admin = _current_admin()
        if admin:
            items = list_noc_requests("it", admin, "Pending")
            pending_noc = len(items or [])
            inventory_items = list_noc_requests("inventory", admin, "Pending")
            pending_inventory_noc = len(inventory_items or [])
    except Exception:
        pending_noc = 0
        pending_inventory_noc = 0

    pending_purchase = 0
    try:
        from .inventory_purchase import ensure_purchase_table
        from .models.inventory_purchase_request import InventoryPurchaseRequest
        from .tenant_context import get_current_tenant_id

        ensure_purchase_table()
        pending_purchase = InventoryPurchaseRequest.query.filter(
            InventoryPurchaseRequest.tenant_id == get_current_tenant_id(),
            InventoryPurchaseRequest.status.in_(("pending_manager", "pending_finance", "finance_approved")),
        ).count()
    except Exception:
        pending_purchase = 0

    pending_requests = (
        pending_returns + pending_emails + pending_assignments + pending_noc + pending_day_use
    )
    return _ok(
        {
            "summary": {
                "total_assets": total_assets,
                "open_tickets": open_tickets,
                "assigned_assets": assigned_assets,
                "day_use_assets": day_use_assets,
                "available_assets": available_assets,
                "available_software": available_software,
                "pending_returns": pending_returns,
                "pending_emails": pending_emails,
                "pending_assignments": pending_assignments,
                "pending_noc": pending_noc,
                "pending_inventory_noc": pending_inventory_noc,
                "pending_day_use": pending_day_use,
                "pending_purchase": pending_purchase,
                "pending_requests": pending_requests,
            }
        }
    )


@it_bp.route("/employees/<string:emp_id>/assets", methods=["GET"])
@jwt_required()
def employee_assets(emp_id):
    emp = _resolve_admin_by_emp_key(emp_id)
    if not emp:
        return _err("Employee not found", 404)

    # Defense-in-depth: non-IT users can only read their own assets.
    claims = get_jwt() or {}
    from .plan_features import can_access_it_panel

    actor = _current_admin()
    if not can_access_it_panel(claims):
        if not actor or int(actor.id) != int(emp.id):
            return _err("You can only view your own assigned assets", 403)

    payload = {"employee": _serialize_admin_employee(emp)}
    return _ok(payload)


@it_bp.route("/employees/<string:emp_id>/return-requests", methods=["GET"])
@jwt_required()
def employee_return_requests(emp_id):
    """Return requests raised by this employee (matches submit flow / requester)."""
    emp = _resolve_admin_by_emp_key(emp_id)
    if not emp:
        return _err("Employee not found", 404)

    # Defense-in-depth: non-IT users can only read their own history.
    claims = get_jwt() or {}
    from .plan_features import can_access_it_panel

    actor = _current_admin()
    if not can_access_it_panel(claims):
        if not actor or int(actor.id) != int(emp.id):
            return _err("You can only view your own return requests", 403)

    status = (request.args.get("status") or "").strip().lower()
    q = ITAssetReturnRequest.query.filter(
        ITAssetReturnRequest.requester_admin_id == emp.id
    )
    if status:
        q = q.filter(ci_equals(ITAssetReturnRequest.status, status))
    rows = _return_requests_newest_first(q)
    return _ok({"requests": [_serialize_return_request(r) for r in rows]})


def _employee_lookup_filter(q, limit=20):
    """Partial match on emp_id, email, or first_name (active employees only)."""
    term = (q or "").strip()
    if len(term) < 2:
        return []
    pattern = f"%{term.lower()}%"
    enabled = db.or_(Admin.is_exited == False, Admin.is_exited.is_(None))
    rows = (
        tenant_admin_query().filter(
            enabled,
            db.or_(
                db.func.lower(db.func.coalesce(Admin.emp_id, "")).like(pattern),
                db.func.lower(db.func.coalesce(Admin.email, "")).like(pattern),
                db.func.lower(db.func.coalesce(Admin.first_name, "")).like(pattern),
            ),
        )
        .order_by(Admin.first_name.asc(), Admin.id.asc())
        .limit(limit)
        .all()
    )
    return rows


@it_bp.route("/employees/lookup", methods=["GET"])
@jwt_required()
def lookup_employee():
    q = (request.args.get("q") or "").strip()
    if not q:
        return _err("q is required")
    if len(q) < 2:
        return _ok({"employees": [], "message": "Type at least 2 characters to search"})

    rows = _employee_lookup_filter(q)
    return _ok({"employees": _serialize_admin_employees(rows)})


def _options_limit():
    limit = request.args.get("limit", 20, type=int) or 20
    return min(max(limit, 1), 50)


@it_bp.route("/employees/options", methods=["GET"])
@jwt_required()
def employee_options():
    """Search-as-you-type picker: active employees, id and name only."""
    enabled = db.or_(Admin.is_exited == False, Admin.is_exited.is_(None))  # noqa: E712
    q = apply_search(tenant_admin_query().filter(enabled), _EMPLOYEE_SEARCH, (request.args.get("q") or "").strip())
    rows = (
        q.options(load_only(Admin.id, Admin.emp_id, Admin.first_name, Admin.email))
        .order_by(Admin.first_name.asc(), Admin.id.asc())
        .limit(_options_limit())
        .all()
    )
    return _ok({
        "options": [
            {"adminId": a.id, "empId": a.emp_id or str(a.id), "name": a.first_name or "", "email": a.email or ""}
            for a in rows
        ]
    })


_INVENTORY_SEARCH = (
    ITInventoryItem.name,
    ITInventoryItem.inventory_code,
    ITInventoryItem.hw_type,
    ITInventoryItem.vendor,
    ITInventoryItem.location,
)
_INVENTORY_SORTS = {
    "name": ITInventoryItem.name,
    "category": ITInventoryItem.category,
    "created_at": ITInventoryItem.created_at,
    "updated_at": ITInventoryItem.updated_at,
    "available": ITInventoryItem.available_quantity,
}


def _inventory_items_query():
    category = (request.args.get("category") or "").strip()
    inventory_category = (request.args.get("inventory_category") or "").strip()
    q = it_query(ITInventoryItem)
    if category:
        q = q.filter(ci_equals(ITInventoryItem.category, category))
    if inventory_category:
        aliases = _inv_cat_aliases(inventory_category)
        if len(aliases) > 1:
            q = q.filter(ITInventoryItem.inventory_category.in_(aliases))
        else:
            q = q.filter(ci_equals(ITInventoryItem.inventory_category, inventory_category))
    return q


@it_bp.route("/inventory/items", methods=["GET"])
@jwt_required()
def list_inventory_items():
    params = list_params()
    order, err = _sort_or_err(params, _INVENTORY_SORTS, ITInventoryItem, (ITInventoryItem.id.desc(),))
    if err:
        return err
    q = apply_search(_inventory_items_query(), _INVENTORY_SEARCH, params.search)
    options = _INVENTORY_MEDIA if params.summary else ()
    items, pagination = list_rows(q, ITInventoryItem, order, params, options)
    return _list_ok("items", [_serialize_inventory_item(i, params.summary) for i in items], pagination)


@it_bp.route("/inventory/items/options", methods=["GET"])
@jwt_required()
def inventory_item_options():
    """Search-as-you-type picker: catalog items, id and name only."""
    q = apply_search(_inventory_items_query(), _INVENTORY_SEARCH, (request.args.get("q") or "").strip())
    rows = (
        q.options(
            load_only(
                ITInventoryItem.id,
                ITInventoryItem.inventory_code,
                ITInventoryItem.name,
                ITInventoryItem.category,
                ITInventoryItem.inventory_category,
            )
        )
        .order_by(ITInventoryItem.name.asc(), ITInventoryItem.id.asc())
        .limit(_options_limit())
        .all()
    )
    return _ok({
        "options": [
            {
                "id": i.id,
                "inventoryCode": i.inventory_code,
                "name": i.name,
                "category": i.category,
                "inventoryCategory": _effective_inventory_category(i),
            }
            for i in rows
        ]
    })


@it_bp.route("/inventory/items", methods=["POST"])
@jwt_required()
def create_inventory_item():
    data = request.get_json(silent=True) or {}
    name = (data.get("name") or "").strip()
    category = (data.get("category") or "").strip()
    if not name or not category:
        return _err("name and category are required")
    inventory_category = _canonical_inv_cat(data.get("inventory_category") or "IT Assets")
    inv_lower = inventory_category.lower()
    if inv_lower in ("office assets", "fixed assets") and category.strip().lower() not in (
        "accessories",
        "consumables",
        "stock",
    ):
        category = "Stock"
    if inv_lower == "infrastructure assets" and category.strip().lower() not in (
        "accessories",
        "consumables",
        "stock",
        "equipment",
    ):
        category = "Stock"
    if inv_lower == "transport assets" and category.strip().lower() != "vehicle":
        category = "Vehicle"

    initial_quantity = int(data.get("initial_quantity") or 0)
    is_qty_managed = _is_qty_managed_inventory(category, inventory_category)
    if is_qty_managed and initial_quantity < 1:
        return _err("initial_quantity must be >= 1 for quantity-based stock items")

    current_admin = _current_admin()
    notes = (data.get("notes") or data.get("remark") or "").strip() or None
    if notes and len(notes) > 500:
        notes = notes[:500]
    photos, photo_err = _stored_media(data.get("photos") or [], "inventory")
    receipts, receipt_err = _stored_media(data.get("receipts") or [], "receipts")
    if photo_err or receipt_err:
        return photo_err or receipt_err
    item = ITInventoryItem(
        inventory_code=data.get("inventory_code") or _next_code("INV", ITInventoryItem, "inventory_code"),
        name=name,
        category=category,
        inventory_category=inventory_category,
        hw_type=(data.get("hw_type") or None),
        photos_json=photos,
        vendor=(data.get("vendor") or "").strip() or None,
        purchase_date=_parse_date(data.get("purchase_date")),
        receipts_json=receipts,
        location=(data.get("location") or "").strip() or None,
        notes=notes,
        total_quantity=initial_quantity if is_qty_managed else 0,
        available_quantity=initial_quantity if is_qty_managed else 0,
        assigned_quantity=0,
        created_by_admin_id=current_admin.id if current_admin else None,
    )
    db.session.add(item)
    db.session.flush()

    from .itam.activity_log_scope import should_log_catalog_receive

    if should_log_catalog_receive(is_qty_managed):
        receive_remark = notes or f"Added {name} to {inventory_category}"
        err = _record_transition_or_err(
            action_code="RECEIVE",
            remark=receive_remark,
            actor_admin_id=current_admin.id if current_admin else None,
            inventory_item_id=item.id,
            from_status=None,
            to_status="available",
            related={"kind": "catalog_create", "inventory_category": inventory_category},
        )
        if err:
            db.session.rollback()
            return err

    db.session.commit()
    return _ok({"item": _serialize_inventory_item(item)}, "Inventory item created", 201)


@it_bp.route("/inventory/items/<int:item_id>", methods=["GET"])
@jwt_required()
def get_inventory_item(item_id):
    item = it_get(ITInventoryItem, item_id)
    if not item:
        return _err("Inventory item not found", 404)
    return _ok({"item": _serialize_inventory_item(item)})


@it_bp.route("/inventory/items/<int:item_id>", methods=["PATCH"])
@jwt_required()
def update_inventory_item(item_id):
    item = it_get(ITInventoryItem, item_id)
    if not item:
        return _err("Inventory item not found", 404)
    data = request.get_json(silent=True) or {}
    for key, attr in (
        ("name", "name"),
        ("category", "category"),
        ("inventory_category", "inventory_category"),
        ("hw_type", "hw_type"),
        ("vendor", "vendor"),
        ("location", "location"),
        ("notes", "notes"),
    ):
        if key in data:
            setattr(item, attr, (data.get(key) or None))
    if "photos" in data:
        photos, err = _stored_media(data.get("photos") or [], "inventory")
        if err:
            return err
        item.photos_json = photos
        from sqlalchemy.orm.attributes import flag_modified
        flag_modified(item, "photos_json")
    if "receipts" in data:
        receipts, err = _stored_media(data.get("receipts") or [], "receipts")
        if err:
            return err
        item.receipts_json = receipts
    if "purchase_date" in data:
        item.purchase_date = _parse_date(data.get("purchase_date"))

    for key, attr in (
        ("total_quantity", "total_quantity"),
        ("available_quantity", "available_quantity"),
        ("assigned_quantity", "assigned_quantity"),
        ("not_working_quantity", "not_working_quantity"),
        ("repair_quantity", "repair_quantity"),
    ):
        if key in data and data.get(key) is not None:
            try:
                setattr(item, attr, max(0, int(data.get(key) or 0)))
            except (TypeError, ValueError):
                return _err(f"{key} must be an integer")
    db.session.commit()
    return _ok({"item": _serialize_inventory_item(item)}, "Inventory item updated")


@it_bp.route("/units", methods=["GET"])
@jwt_required()
def list_units():
    params = list_params()
    order, err = _sort_or_err(params, _UNIT_SORTS, ITAssetUnit, (ITAssetUnit.id.desc(),))
    if err:
        return err
    q = apply_search(_units_query(pagination_enabled()), _UNIT_SEARCH, params.search)
    rows, pagination = units_for_listing(q, order, params)
    return _list_ok("units", serialize_listed_units(rows, params.summary), pagination)


def serialize_listed_units(rows, summary=False):
    """Unit list payload, with the latest transition attached when the timeline is on."""
    from .itam.timeline_service import latest_by_unit_ids, timeline_enabled

    payload = [_serialize_asset_unit(r, summary) for r in rows]
    if timeline_enabled(current_app.config) and payload:
        latest_map = latest_by_unit_ids([r.id for r in rows])
        for item in payload:
            latest = latest_map.get(int(item["id"]))
            if latest:
                item["lastTransition"] = latest
                item["lastRemark"] = latest.get("remark")
                item["lastTransitionAt"] = latest.get("occurredAt")
    return payload


_UNIT_SORTS = {
    "asset_name": ITAssetUnit.asset_name,
    "unit_code": ITAssetUnit.unit_code,
    "serial_number": ITAssetUnit.serial_number,
    "status": ITAssetUnit.status,
    "hw_type": ITAssetUnit.hw_type,
    "assigned_at": ITAssetUnit.assigned_at,
    "created_at": ITAssetUnit.created_at,
    "updated_at": ITAssetUnit.updated_at,
}


def _units_query(extra_filters):
    """Units for this tenant filtered by status / inventory_id, plus hw_type / category when ``extra_filters``."""
    status = (request.args.get("status") or "").strip()
    inventory_id = request.args.get("inventory_id", type=int)
    q = it_query(ITAssetUnit)
    if status:
        q = q.filter(ci_equals(ITAssetUnit.status, status))
    if inventory_id:
        q = q.filter(ITAssetUnit.inventory_item_id == inventory_id)
    if extra_filters:
        hw_type = (request.args.get("hw_type") or "").strip()
        category = (request.args.get("category") or "").strip()
        if hw_type:
            q = q.filter(ci_equals(ITAssetUnit.hw_type, hw_type))
        if category:
            q = q.filter(ci_equals(ITAssetUnit.category, category))
    return q


@it_bp.route("/units/options", methods=["GET"])
@jwt_required()
def unit_options():
    """Search-as-you-type picker: units, id and labels only."""
    q = apply_search(_units_query(True), _UNIT_SEARCH, (request.args.get("q") or "").strip())
    rows = (
        q.options(
            load_only(
                ITAssetUnit.id,
                ITAssetUnit.unit_code,
                ITAssetUnit.asset_name,
                ITAssetUnit.hw_type,
                ITAssetUnit.serial_number,
                ITAssetUnit.status,
            )
        )
        .order_by(ITAssetUnit.asset_name.asc(), ITAssetUnit.id.asc())
        .limit(_options_limit())
        .all()
    )
    return _ok({
        "options": [
            {
                "id": u.id,
                "unitCode": u.unit_code,
                "assetName": u.asset_name,
                "hwType": u.hw_type,
                "serialNumber": u.serial_number,
                "status": u.status,
            }
            for u in rows
        ]
    })


DEFAULT_DEVICE_TYPES = ("EVT", "DVT", "PVT", "DV")


def _normalize_device_type(value):
    text = str(value or "").strip()
    if not text:
        return None
    return text[:40]


def _remember_device_type(name):
    """Persist a custom / selected device type so later dropdowns can reuse it."""
    cleaned = _normalize_device_type(name)
    if not cleaned:
        return None
    existing = (
        it_query(ITDeviceType)
        .filter(db.func.lower(ITDeviceType.name) == cleaned.casefold())
        .first()
    )
    if existing:
        return existing
    row = ITDeviceType(tenant_id=get_current_tenant_id(), name=cleaned)
    db.session.add(row)
    return row


def _device_type_names():
    names = list(DEFAULT_DEVICE_TYPES)
    seen = {n.casefold() for n in names}
    catalog = it_query(ITDeviceType).order_by(ITDeviceType.name.asc()).all()
    for row in catalog:
        label = (row.name or "").strip()
        if label and label.casefold() not in seen:
            names.append(label)
            seen.add(label.casefold())
    used = (
        it_query(ITAssetUnit)
        .with_entities(ITAssetUnit.device_type)
        .filter(ITAssetUnit.device_type.isnot(None))
        .distinct()
        .all()
    )
    for (label,) in used:
        text = (label or "").strip()
        if text and text.casefold() not in seen:
            names.append(text)
            seen.add(text.casefold())
    return names


@it_bp.route("/device-types", methods=["GET"])
@jwt_required()
def list_device_types():
    return _ok({"deviceTypes": _device_type_names()})


@it_bp.route("/device-types", methods=["POST"])
@jwt_required()
def create_device_type():
    data = request.get_json(silent=True) or {}
    name = _normalize_device_type(data.get("name") or data.get("device_type"))
    if not name:
        return _err("Device type name is required")
    _remember_device_type(name)
    db.session.commit()
    return _ok({"deviceTypes": _device_type_names(), "name": name}, "Device type saved", 201)


def _is_mobile_tablet_hw(hw_type) -> bool:
    return str(hw_type or "").strip().lower() in ("mobile", "tablet")


def _new_asset_unit_code(hw_type=None) -> str:
    prefix = (str(hw_type or "UNIT").strip() or "UNIT").upper().replace(" ", "")[:10]
    return f"{prefix}-{uuid.uuid4().hex[:12]}"


@it_bp.route("/units/bulk", methods=["POST"])
@jwt_required()
def create_units_bulk():
    data = request.get_json(silent=True) or {}
    inventory_id = data.get("inventory_item_id")
    units = data.get("units") or []
    if not inventory_id or not isinstance(units, list) or not units:
        return _err("inventory_item_id and non-empty units array are required")
    inv = it_get(ITInventoryItem, inventory_id)
    if not inv:
        return _err("Inventory item not found", 404)

    prepared = []
    seen_codes = set()
    dup_in_request = []
    for row in units:
        hw_type = row.get("hw_type") or inv.hw_type
        serial = (row.get("serial_number") or "").strip()
        if _is_mobile_tablet_hw(hw_type):
            unit_code = _new_asset_unit_code(hw_type)
            prepared.append((unit_code, row))
            continue
        unit_code = (row.get("unit_code") or serial or "").strip()
        if not unit_code:
            continue
        code_key = unit_code.casefold()
        if code_key in seen_codes:
            if unit_code not in dup_in_request:
                dup_in_request.append(unit_code)
            continue
        seen_codes.add(code_key)
        prepared.append((unit_code, row))

    if dup_in_request:
        listed = ", ".join(dup_in_request[:10])
        return _err(
            f"Duplicate serial / unit code(s) in this request: {listed}. "
            "Each asset must have a unique serial number.",
            409,
        )
    if not prepared:
        return _err("No valid units provided")

    laptop_codes = [c for c, r in prepared if not _is_mobile_tablet_hw(r.get("hw_type") or inv.hw_type)]
    if laptop_codes:
        existing = (
            it_query(ITAssetUnit).filter(
                db.func.lower(ITAssetUnit.unit_code).in_([c.casefold() for c in laptop_codes])
            )
            .with_entities(ITAssetUnit.unit_code)
            .all()
        )
        if existing:
            codes = sorted({row[0] for row in existing if row and row[0]})
            listed = ", ".join(codes[:10])
            more = f" (+{len(codes) - 10} more)" if len(codes) > 10 else ""
            return _err(
                f"Asset unit code(s) already exist: {listed}{more}. "
                "Use a unique serial / unit code for each device.",
                409,
            )

    created = []
    for unit_code, row in prepared:
        photos, err = _stored_media(row.get("photos") or [], "units")
        if err:
            return err
        unit = ITAssetUnit(
            tenant_id=int(inv.tenant_id) if inv.tenant_id is not None else get_current_tenant_id(),
            unit_code=unit_code,
            inventory_item_id=inventory_id,
            asset_name=(row.get("asset_name") or inv.name or "").strip(),
            category=(row.get("category") or inv.category or "Hardware").strip(),
            hw_type=row.get("hw_type") or inv.hw_type,
            brand=row.get("brand"),
            make=row.get("make"),
            model=row.get("model"),
            serial_number=row.get("serial_number"),
            device_type=_normalize_device_type(row.get("device_type") or row.get("deviceType")),
            imei1=row.get("imei1"),
            imei2=row.get("imei2"),
            project_code=(row.get("project_code") or row.get("projectCode") or "").strip() or None,
            device_location=(row.get("device_location") or row.get("deviceLocation") or "").strip() or None,
            status=(row.get("status") or "available").strip(),
            photos_json=photos,
        )
        unit_notes = (row.get("notes") or row.get("remark") or "").strip() or None
        if unit_notes and len(unit_notes) > 500:
            unit_notes = unit_notes[:500]
        if unit.device_type:
            _remember_device_type(unit.device_type)
        db.session.add(unit)
        created.append((unit, unit_notes))

    batch_notes = (data.get("notes") or data.get("remark") or "").strip() or None
    if batch_notes and len(batch_notes) > 500:
        batch_notes = batch_notes[:500]
    first_unit_notes = next((n for _, n in created if n), None)
    if not (inv.notes or "").strip():
        inv.notes = first_unit_notes or batch_notes

    try:
        db.session.flush()
        actor = _current_admin()
        for unit, unit_notes in created:
            remark = unit_notes or batch_notes
            err = _record_transition_or_err(
                action_code="RECEIVE",
                remark=remark,
                actor_admin_id=actor.id if actor else None,
                asset_unit_id=unit.id,
                inventory_item_id=inventory_id,
                from_status=None,
                to_status=unit.status or "available",
            )
            if err:
                db.session.rollback()
                return err
        _recalc_inventory_counts(inventory_id)
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        from sqlalchemy.exc import IntegrityError

        if isinstance(e, IntegrityError):
            return _err(
                "Could not save one or more assets. Please try again.",
                409,
            )
        raise
    return _ok(
        {"units": [_serialize_asset_unit(u) for u, _ in created]},
        "Units created",
        201,
    )


@it_bp.route("/units/<int:unit_id>/status", methods=["PATCH"])
@jwt_required()
def update_unit_status(unit_id):
    from .itam.transition_service import action_for_unit_status, extract_remark_fields

    unit = it_get(ITAssetUnit, unit_id)
    if not unit:
        return _err("Unit not found", 404)
    data = request.get_json(silent=True) or {}
    status = (data.get("status") or "").strip()
    if not status:
        return _err("status is required")

    fields = extract_remark_fields(data)
    from_status = unit.status
    action = action_for_unit_status(status, from_status=from_status)
    if action == "COMPLETE_REPAIR" and not fields.get("condition_grade"):
        fields["condition_grade"] = (data.get("condition_grade") or "B").strip() or "B"

    actor = _current_admin()
    err = _record_transition_or_err(
        action_code=action,
        remark=fields.get("remark"),
        reason_code=fields.get("reason_code"),
        condition_grade=fields.get("condition_grade"),
        actor_admin_id=actor.id if actor else None,
        asset_unit_id=unit.id,
        inventory_item_id=unit.inventory_item_id,
        from_status=from_status,
        to_status=status,
    )
    if err:
        return err

    unit.status = status
    if status == "repair" and not unit.repair_date:
        unit.repair_date = utc_now()
    if status == "available":
        unit.repair_date = None
        unit.assigned_to_admin_id = None
        unit.assigned_at = None

    from .itam.lifecycle_service import lifecycle_for_legacy_status_change

    life, ctype = lifecycle_for_legacy_status_change(status, from_unit=unit)
    custody = {"type": ctype}
    if ctype == "VENDOR":
        custody["vendor_name"] = (data.get("vendor_name") or data.get("vendor") or "").strip() or None
    err = _apply_lifecycle_or_err(
        unit,
        lifecycle_status=life,
        custody_type=ctype,
        custody=custody,
        actor_admin_id=actor.id if actor else None,
        dual_write_legacy=False,  # caller already set legacy status
    )
    if err:
        return err

    _recalc_inventory_counts(unit.inventory_item_id)
    db.session.commit()
    return _ok({"unit": _serialize_asset_unit(unit)}, "Unit status updated")


@it_bp.route("/units/<int:unit_id>", methods=["PATCH"])
@jwt_required()
def update_unit(unit_id):
    """Update identity fields on a unit. Status / assignment stay on their own endpoints."""
    unit = it_get(ITAssetUnit, unit_id)
    if not unit:
        return _err("Unit not found", 404)
    data = request.get_json(silent=True) or {}
    for key, attr in (
        ("asset_name", "asset_name"),
        ("brand", "brand"),
        ("make", "make"),
        ("model", "model"),
        ("serial_number", "serial_number"),
        ("device_type", "device_type"),
        ("deviceType", "device_type"),
        ("imei1", "imei1"),
        ("imei2", "imei2"),
        ("project_code", "project_code"),
        ("device_location", "device_location"),
    ):
        if key in data:
            raw = data.get(key)
            value = None if raw is None else str(raw).strip()
            setattr(unit, attr, value or None)
    if "photos" in data:
        photos, err = _stored_media(data.get("photos") or [], "units")
        if err:
            return err
        unit.photos_json = photos
        from sqlalchemy.orm.attributes import flag_modified
        flag_modified(unit, "photos_json")
    if not (unit.asset_name or "").strip():
        return _err("asset_name is required")
    if "device_type" in data or "deviceType" in data:
        unit.device_type = _normalize_device_type(unit.device_type)
        if unit.device_type:
            _remember_device_type(unit.device_type)
    db.session.commit()
    return _ok({"unit": _serialize_asset_unit(unit)}, "Unit updated")


@it_bp.route("/units/<int:unit_id>", methods=["DELETE"])
@jwt_required()
def delete_unit(unit_id):
    from .itam.transition_service import extract_remark_fields

    unit = it_get(ITAssetUnit, unit_id)
    if not unit:
        return _err("Unit not found", 404)
    if unit.status == "assigned":
        return _err("Unassign this unit before permanent removal", 400)

    data = request.get_json(silent=True) or {}
    fields = extract_remark_fields(data)
    actor = _current_admin()
    inv_id = unit.inventory_item_id
    from_status = unit.status

    err = _record_transition_or_err(
        action_code="RETIRE",
        remark=fields.get("remark"),
        reason_code=fields.get("reason_code") or "PERMANENT_DELETE",
        condition_grade=fields.get("condition_grade") or "Fail",
        actor_admin_id=actor.id if actor else None,
        asset_unit_id=unit.id,
        inventory_item_id=inv_id,
        from_status=from_status,
        to_status="Retired",
    )
    if err:
        return err

    from flask import current_app
    from .itam.lifecycle_service import close_open_custodies, lifecycle_enabled

    if lifecycle_enabled(current_app.config):
        close_open_custodies(asset_unit_id=unit.id)
        err = _apply_lifecycle_or_err(
            unit,
            lifecycle_status="Retired",
            custody_type="NONE",
            custody={"type": "NONE"},
            actor_admin_id=actor.id if actor else None,
            dual_write_legacy=False,
        )
        if err:
            return err

    ITDeletedAssetLog.query.filter_by(asset_unit_id=unit_id).update(
        {ITDeletedAssetLog.asset_unit_id: None},
        synchronize_session=False,
    )
    ITRemovedAsset.query.filter_by(asset_unit_id=unit_id).update(
        {ITRemovedAsset.asset_unit_id: None},
        synchronize_session=False,
    )
    ITParcelExportItem.query.filter_by(asset_unit_id=unit_id).update(
        {ITParcelExportItem.asset_unit_id: None},
        synchronize_session=False,
    )

    db.session.delete(unit)
    _recalc_inventory_counts(inv_id)
    db.session.commit()
    return _ok(message="Unit deleted")


@it_bp.route("/assignments/units", methods=["POST"])
@jwt_required()
def assign_unit():
    data = request.get_json(silent=True) or {}
    unit_id = data.get("unit_id")
    assigned_to_admin_id = data.get("assigned_to_admin_id")
    assigned_to_emp_id = (data.get("assigned_to_emp_id") or "").strip()
    if not unit_id or (not assigned_to_admin_id and not assigned_to_emp_id):
        return _err("unit_id and assigned_to_admin_id/assigned_to_emp_id are required")

    unit = it_get(ITAssetUnit, unit_id)
    target_admin = get_admin_in_tenant(assigned_to_admin_id) if assigned_to_admin_id else None
    if not target_admin and assigned_to_emp_id:
        target_admin = tenant_admin_query().filter(
            db.func.lower(db.func.coalesce(Admin.emp_id, "")) == assigned_to_emp_id.lower()
        ).first()
    actor = _current_admin()
    if not unit:
        return _err("Unit not found", 404)
    if not target_admin:
        return _err("Target admin not found", 404)
    if unit.status == "assigned":
        return _err("Unit is already assigned")
    if (unit.status or "").lower() == "daily_out":
        return _err("Unit is on loan checkout")
    from .models.daily_checkout import DailyUnitHold

    if DailyUnitHold.query.get(unit.id):
        return _err("Unit is on loan checkout")

    from .itam.transition_service import extract_remark_fields

    fields = extract_remark_fields(data)
    err = _record_transition_or_err(
        action_code="CHECKOUT",
        remark=fields.get("remark"),
        reason_code=fields.get("reason_code"),
        condition_grade=fields.get("condition_grade"),
        actor_admin_id=actor.id if actor else None,
        asset_unit_id=unit.id,
        inventory_item_id=unit.inventory_item_id,
        from_status=unit.status,
        to_status="assigned",
        from_custody={"type": "NONE"},
        to_custody={
            "type": "EMPLOYEE",
            "admin_id": target_admin.id,
            "emp_id": getattr(target_admin, "emp_id", None),
        },
    )
    if err:
        return err

    unit.status = "assigned"
    unit.assigned_to_admin_id = target_admin.id
    unit.assigned_at = _parse_dt(data.get("assigned_at")) or utc_now()
    unit.asset_tag = (data.get("asset_tag") or unit.asset_tag)
    photos, err = _stored_media(data.get("assignment_photos") or [], "assignments")
    if data.get("assignment_photos") and err:
        return err
    unit.assignment_photos_json = photos or unit.assignment_photos_json or []

    err = _apply_lifecycle_or_err(
        unit,
        lifecycle_status="CheckedOut",
        custody_type="EMPLOYEE",
        custody={
            "type": "EMPLOYEE",
            "admin_id": target_admin.id,
            "emp_id": getattr(target_admin, "emp_id", None),
        },
        actor_admin_id=actor.id if actor else None,
        dual_write_legacy=False,
    )
    if err:
        return err

    assignment = ITAssetAssignment(
        assignment_type="assign",
        assigned_to_admin_id=target_admin.id,
        assigned_by_admin_id=actor.id if actor else None,
        asset_unit_id=unit.id,
        notes=fields.get("remark") or data.get("notes"),
        assigned_at=unit.assigned_at,
    )
    db.session.add(assignment)
    _recalc_inventory_counts(unit.inventory_item_id)
    db.session.commit()
    try:
        ok, msg = send_it_assignment_notification(
            target_admin=target_admin,
            actor_admin=actor,
            assignment_kind="unit",
            unit=unit,
        )
        if ok:
            current_app.logger.info(
                "[IT assign_unit] notification sent | unit_id=%s target_admin_id=%s msg=%s",
                unit.id,
                target_admin.id,
                msg,
            )
        else:
            current_app.logger.warning(
                "[IT assign_unit] notification failed | unit_id=%s target_admin_id=%s msg=%s",
                unit.id,
                target_admin.id,
                msg,
            )
    except Exception as e:
        current_app.logger.warning(
            "[IT assign_unit] notification exception | unit_id=%s target_admin_id=%s err=%s",
            unit.id,
            target_admin.id,
            e,
        )
    return _ok({"unit": _serialize_asset_unit(unit), "assignment_id": assignment.id}, "Unit assigned")


@it_bp.route("/assignments/units/<int:unit_id>/return", methods=["POST"])
@jwt_required()
def return_unit(unit_id):
    from .itam.transition_service import (
        action_for_return_destination,
        extract_remark_fields,
    )

    data = request.get_json(silent=True) or {}
    unit = it_get(ITAssetUnit, unit_id)
    actor = _current_admin()
    if not unit:
        return _err("Unit not found", 404)

    previous_assigned_to = unit.assigned_to_admin_id
    if not previous_assigned_to:
        return _err("Unit is not assigned")

    new_status = (data.get("status") or "available").strip()
    fields = extract_remark_fields(data)
    action = action_for_return_destination(new_status)
    if action == "CHECKIN" and not fields.get("condition_grade"):
        fields["condition_grade"] = (data.get("condition_grade") or "B").strip() or "B"

    err = _record_transition_or_err(
        action_code=action,
        remark=fields.get("remark"),
        reason_code=fields.get("reason_code"),
        condition_grade=fields.get("condition_grade"),
        actor_admin_id=actor.id if actor else None,
        asset_unit_id=unit.id,
        inventory_item_id=unit.inventory_item_id,
        from_status=unit.status,
        to_status=new_status,
        from_custody={"type": "EMPLOYEE", "admin_id": previous_assigned_to},
        to_custody={"type": "NONE"},
    )
    if err:
        return err

    unit.status = new_status
    unit.assigned_to_admin_id = None
    unit.assigned_at = None
    if new_status == "repair":
        unit.repair_date = utc_now()
    elif new_status == "available":
        unit.repair_date = None

    from .itam.lifecycle_service import lifecycle_for_legacy_status_change

    life, ctype = lifecycle_for_legacy_status_change(new_status)
    err = _apply_lifecycle_or_err(
        unit,
        lifecycle_status=life,
        custody_type=ctype,
        custody={"type": ctype},
        actor_admin_id=actor.id if actor else None,
        dual_write_legacy=False,
    )
    if err:
        return err

    assignment = ITAssetAssignment(
        assignment_type="return",
        assigned_to_admin_id=previous_assigned_to,
        assigned_by_admin_id=actor.id if actor else None,
        asset_unit_id=unit.id,
        notes=fields.get("remark") or data.get("notes"),
        assigned_at=utc_now(),
        unassigned_at=utc_now(),
    )
    db.session.add(assignment)
    _recalc_inventory_counts(unit.inventory_item_id)
    db.session.commit()
    return _ok({"unit": _serialize_asset_unit(unit), "assignment_id": assignment.id}, "Unit returned")


@it_bp.route("/software/licenses", methods=["GET"])
@jwt_required()
def list_software_licenses():
    name = (request.args.get("name") or "").strip()
    status = (request.args.get("status") or "").strip()
    q = it_query(ITSoftwareLicense)
    if name:
        q = q.filter(ci_equals(ITSoftwareLicense.name, name))
    if status:
        q = q.filter(ci_equals(ITSoftwareLicense.status, status))
    params = list_params()
    order, err = _sort_or_err(params, _LICENSE_SORTS, ITSoftwareLicense, (ITSoftwareLicense.id.desc(),))
    if err:
        return err
    q = apply_search(q, (ITSoftwareLicense.name, ITSoftwareLicense.license_code), params.search)
    rows, pagination = list_rows(
        q,
        ITSoftwareLicense,
        order,
        params,
        (selectinload(ITSoftwareLicense.assigned_to_admin),),
        sort_ids_first=False,
    )
    prime_photo_urls(r.assigned_to_admin_id for r in rows)
    return _list_ok("licenses", [_serialize_license(r) for r in rows], pagination)


_LICENSE_SORTS = {
    "name": ITSoftwareLicense.name,
    "status": ITSoftwareLicense.status,
    "subscription_end": ITSoftwareLicense.subscription_end,
    "assigned_at": ITSoftwareLicense.assigned_at,
    "created_at": ITSoftwareLicense.created_at,
}


@it_bp.route("/software/licenses/bulk", methods=["POST"])
@jwt_required()
def create_software_licenses():
    data = request.get_json(silent=True) or {}
    inventory_id = data.get("inventory_item_id")
    name = (data.get("name") or "").strip()
    quantity = int(data.get("quantity") or 0)
    if not inventory_id or not name or quantity < 1:
        return _err("inventory_item_id, name and quantity>=1 are required")
    inv = it_get(ITInventoryItem, inventory_id)
    if not inv:
        return _err("Inventory item not found", 404)

    created = []
    start_date = _parse_date(data.get("subscription_start"))
    end_date = _parse_date(data.get("subscription_end"))
    batch_notes = (data.get("notes") or data.get("remark") or "").strip() or None
    if batch_notes and len(batch_notes) > 500:
        batch_notes = batch_notes[:500]
    for _ in range(quantity):
        lic = ITSoftwareLicense(
            tenant_id=int(inv.tenant_id) if inv.tenant_id is not None else get_current_tenant_id(),
            license_code=_next_code("SW", ITSoftwareLicense, "license_code"),
            inventory_item_id=inventory_id,
            name=name,
            subscription_start=start_date,
            subscription_end=end_date,
            status="available",
        )
        db.session.add(lic)
        db.session.flush()
        created.append(lic)
    db.session.flush()
    actor = _current_admin()
    for lic in created:
        err = _record_transition_or_err(
            action_code="RECEIVE",
            remark=batch_notes,
            actor_admin_id=actor.id if actor else None,
            software_license_id=lic.id,
            inventory_item_id=inventory_id,
            from_status=None,
            to_status=lic.status or "available",
        )
        if err:
            db.session.rollback()
            return err
    _recalc_inventory_counts(inventory_id)
    db.session.commit()
    return _ok({"licenses": [_serialize_license(i) for i in created]}, "Software licenses created", 201)


@it_bp.route("/assignments/software", methods=["POST"])
@jwt_required()
def assign_software():
    data = request.get_json(silent=True) or {}
    license_id = data.get("license_id")
    assigned_to_admin_id = data.get("assigned_to_admin_id")
    assigned_to_emp_id = (data.get("assigned_to_emp_id") or "").strip()
    if not license_id or (not assigned_to_admin_id and not assigned_to_emp_id):
        return _err("license_id and assigned_to_admin_id/assigned_to_emp_id are required")
    lic = it_get(ITSoftwareLicense, license_id)
    target_admin = get_admin_in_tenant(assigned_to_admin_id) if assigned_to_admin_id else None
    if not target_admin and assigned_to_emp_id:
        target_admin = tenant_admin_query().filter(
            db.func.lower(db.func.coalesce(Admin.emp_id, "")) == assigned_to_emp_id.lower()
        ).first()
    actor = _current_admin()
    if not lic:
        return _err("Software license not found", 404)
    if not target_admin:
        return _err("Target admin not found", 404)
    if lic.status == "assigned":
        return _err("License is already assigned")

    from .itam.transition_service import extract_remark_fields

    fields = extract_remark_fields(data)
    err = _record_transition_or_err(
        action_code="CHECKOUT",
        remark=fields.get("remark"),
        reason_code=fields.get("reason_code"),
        condition_grade=fields.get("condition_grade"),
        actor_admin_id=actor.id if actor else None,
        software_license_id=lic.id,
        inventory_item_id=lic.inventory_item_id,
        from_status=lic.status,
        to_status="assigned",
        to_custody={
            "type": "EMPLOYEE",
            "admin_id": target_admin.id,
            "emp_id": getattr(target_admin, "emp_id", None),
        },
    )
    if err:
        return err

    lic.status = "assigned"
    lic.assigned_to_admin_id = target_admin.id
    lic.assigned_at = utc_now()

    assignment = ITAssetAssignment(
        assignment_type="assign",
        assigned_to_admin_id=target_admin.id,
        assigned_by_admin_id=actor.id if actor else None,
        software_license_id=lic.id,
        notes=fields.get("remark") or data.get("notes"),
    )
    db.session.add(assignment)
    _recalc_inventory_counts(lic.inventory_item_id)
    db.session.commit()
    try:
        ok, msg = send_it_assignment_notification(
            target_admin=target_admin,
            actor_admin=actor,
            assignment_kind="software",
            license_obj=lic,
        )
        if ok:
            current_app.logger.info(
                "[IT assign_software] notification sent | license_id=%s target_admin_id=%s msg=%s",
                lic.id,
                target_admin.id,
                msg,
            )
        else:
            current_app.logger.warning(
                "[IT assign_software] notification failed | license_id=%s target_admin_id=%s msg=%s",
                lic.id,
                target_admin.id,
                msg,
            )
    except Exception as e:
        current_app.logger.warning(
            "[IT assign_software] notification exception | license_id=%s target_admin_id=%s err=%s",
            lic.id,
            target_admin.id,
            e,
        )
    return _ok({"license": _serialize_license(lic), "assignment_id": assignment.id}, "Software assigned")


@it_bp.route("/assignments/software/<int:license_id>/return", methods=["POST"])
@jwt_required()
def return_software(license_id):
    from .itam.transition_service import extract_remark_fields

    data = request.get_json(silent=True) or {}
    lic = it_get(ITSoftwareLicense, license_id)
    actor = _current_admin()
    if not lic:
        return _err("Software license not found", 404)
    if lic.status != "assigned":
        return _err("Software license is not assigned")

    fields = extract_remark_fields(data)
    if not fields.get("condition_grade"):
        fields["condition_grade"] = (data.get("condition_grade") or "B").strip() or "B"
    prev_target = lic.assigned_to_admin_id
    err = _record_transition_or_err(
        action_code="CHECKIN",
        remark=fields.get("remark"),
        reason_code=fields.get("reason_code"),
        condition_grade=fields.get("condition_grade"),
        actor_admin_id=actor.id if actor else None,
        software_license_id=lic.id,
        inventory_item_id=lic.inventory_item_id,
        from_status=lic.status,
        to_status="available",
        from_custody={"type": "EMPLOYEE", "admin_id": prev_target},
        to_custody={"type": "NONE"},
    )
    if err:
        return err

    lic.status = "available"
    lic.assigned_to_admin_id = None
    lic.assigned_at = None

    assignment = ITAssetAssignment(
        assignment_type="return",
        assigned_to_admin_id=prev_target,
        assigned_by_admin_id=actor.id if actor else None,
        software_license_id=lic.id,
        notes=fields.get("remark") or data.get("notes"),
        unassigned_at=utc_now(),
    )
    db.session.add(assignment)
    _recalc_inventory_counts(lic.inventory_item_id)
    db.session.commit()
    return _ok({"license": _serialize_license(lic), "assignment_id": assignment.id}, "Software returned")


@it_bp.route("/return-requests", methods=["POST"])
@jwt_required()
def create_return_request():
    actor = _current_admin()
    if not actor:
        return _err("Unauthorized", 401)
    data = request.get_json(silent=True) or {}
    reason = (data.get("reason") or "").strip()
    unit_id = data.get("asset_unit_id")
    license_id = data.get("software_license_id")
    inventory_id = data.get("inventory_item_id")
    quantity = int(data.get("quantity") or 1)
    return_destination = (data.get("return_destination") or "available").strip().lower()
    photos, err = _stored_media(data.get("photos") or [], "returns")
    if err:
        return err
    if not reason:
        return _err("reason is required")
    if return_destination not in ("available", "removed_from_it"):
        return _err("return_destination must be available or removed_from_it")
    if not any([unit_id, license_id, inventory_id]):
        return _err("asset_unit_id/software_license_id/inventory_item_id is required")

    asset_name = None
    category = None
    if unit_id:
        unit = it_get(ITAssetUnit, unit_id)
        if not unit:
            return _err("Unit not found", 404)
        if unit.assigned_to_admin_id != actor.id:
            return _err("You can only request return for your assigned assets", 403)
        asset_name = unit.asset_name or unit.brand or unit.unit_code
        category = unit.category or "Hardware"
    elif license_id:
        lic = it_get(ITSoftwareLicense, license_id)
        if not lic:
            return _err("Software license not found", 404)
        if lic.assigned_to_admin_id != actor.id:
            return _err("You can only request return for your assigned licenses", 403)
        asset_name = lic.name or lic.license_code
        category = "Software"
    else:
        item = it_get(ITInventoryItem, inventory_id)
        if not item:
            return _err("Inventory item not found", 404)
        if quantity < 1:
            return _err("quantity must be >= 1")
        qa = ITInventoryQuantityAssignment.query.filter_by(
            inventory_item_id=item.id,
            assigned_to_admin_id=actor.id,
        ).first()
        assigned_qty = int(qa.quantity or 0) if qa else 0
        if assigned_qty < quantity:
            return _err("You do not have enough assigned quantity to return", 403)
        asset_name = item.name
        category = item.category

    dup = ITAssetReturnRequest.query.filter(
        ITAssetReturnRequest.status.in_(["pending", "approved"]),
        ITAssetReturnRequest.requester_admin_id == actor.id,
        ITAssetReturnRequest.asset_unit_id == unit_id,
        ITAssetReturnRequest.software_license_id == license_id,
        ITAssetReturnRequest.inventory_item_id == inventory_id,
    ).first()
    if dup:
        return _err("A return request for this asset is already pending/approved", 409)

    err = _record_transition_or_err(
        action_code="REQUEST_RETURN",
        remark=reason,
        actor_admin_id=actor.id,
        asset_unit_id=unit_id,
        software_license_id=license_id,
        inventory_item_id=inventory_id,
        from_status="assigned",
        to_status="PendingReturn",
        related={"return_destination": return_destination},
    )
    if err:
        return err

    req = ITAssetReturnRequest(
        request_code=_next_code("RTR", ITAssetReturnRequest, "request_code"),
        requester_admin_id=actor.id,
        requester_emp_id=(actor.emp_id or str(actor.id)),
        asset_unit_id=unit_id,
        software_license_id=license_id,
        inventory_item_id=inventory_id,
        quantity=quantity,
        asset_name=asset_name,
        category=category,
        reason=reason,
        return_destination=return_destination,
        photos_json=photos,
        status="pending",
    )
    db.session.add(req)
    db.session.commit()

    dest_label = (
        "Removed From IT"
        if return_destination == "removed_from_it"
        else "Available"
    )
    ok, msg = send_it_return_request_email(
        requester_admin=actor,
        reason=f"{reason} | Return to: {dest_label}",
        asset_label=f"{asset_name or '-'} ({category or '-'})",
    )
    if ok:
        current_app.logger.info("[IT return-request] mail sent | request_id=%s msg=%s", req.id, msg)
    else:
        current_app.logger.warning("[IT return-request] mail failed | request_id=%s msg=%s", req.id, msg)
    return _ok({"request": _serialize_return_request(req)}, "Return request created", 201)


@it_bp.route("/return-requests", methods=["GET"])
@jwt_required()
def list_return_requests():
    status = (request.args.get("status") or "").strip().lower()
    only_mine = str(request.args.get("mine") or "").strip().lower() in {"1", "true", "yes"}
    emp_id = (request.args.get("emp_id") or "").strip()
    actor = _current_admin()
    params = list_params()
    order, err = _sort_or_err(params, _RETURN_REQUEST_SORTS, ITAssetReturnRequest, _RETURN_REQUESTS_NEWEST)
    if err:
        return err
    q = it_query(ITAssetReturnRequest)
    if status:
        q = q.filter(ci_equals(ITAssetReturnRequest.status, status))
    if emp_id:
        emp = _resolve_admin_by_emp_key(emp_id)
        if not emp:
            return _list_ok("requests", [], empty_pagination(params))
        q = q.filter(ITAssetReturnRequest.requester_admin_id == emp.id)
    elif only_mine and actor:
        q = q.filter(ITAssetReturnRequest.requester_admin_id == actor.id)
    q = apply_search(q, _RETURN_REQUEST_SEARCH, params.search)
    options = _RETURN_REQUEST_ADMINS + ((defer(ITAssetReturnRequest.photos_json),) if params.summary else ())
    rows, pagination = list_rows(q, ITAssetReturnRequest, order, params, options)
    return _list_ok("requests", [_serialize_return_request(r, params.summary) for r in rows], pagination)


_RETURN_REQUEST_SEARCH = (
    ITAssetReturnRequest.request_code,
    ITAssetReturnRequest.asset_name,
    ITAssetReturnRequest.requester_emp_id,
    ITAssetReturnRequest.reason,
)
_RETURN_REQUEST_SORTS = {
    "created_at": ITAssetReturnRequest.created_at,
    "updated_at": ITAssetReturnRequest.updated_at,
    "status": ITAssetReturnRequest.status,
    "asset_name": ITAssetReturnRequest.asset_name,
}


@it_bp.route("/return-requests/<int:request_id>/approve", methods=["PATCH"])
@jwt_required()
def approve_return_request(request_id):
    from .itam.transition_service import extract_remark_fields

    actor = _current_admin()
    row = it_get(ITAssetReturnRequest, request_id)
    if not row:
        return _err("Return request not found", 404)
    if row.status != "pending":
        return _err("Only pending requests can be approved")
    data = request.get_json(silent=True) or {}
    fields = extract_remark_fields(data)
    if not fields.get("remark"):
        fields["remark"] = f"Approved return request {row.request_code}"
    err = _record_transition_or_err(
        action_code="APPROVE_RETURN",
        remark=fields.get("remark"),
        reason_code=fields.get("reason_code"),
        condition_grade=fields.get("condition_grade"),
        actor_admin_id=actor.id if actor else None,
        asset_unit_id=row.asset_unit_id,
        software_license_id=row.software_license_id,
        inventory_item_id=row.inventory_item_id,
        from_status="pending",
        to_status="approved",
        related={"return_request_id": row.id, "request_code": row.request_code},
    )
    if err:
        return err
    row.status = "approved"
    row.approved_by_admin_id = actor.id if actor else None
    row.approved_at = utc_now()
    db.session.commit()
    if row.requester_admin:
        send_it_return_request_status_email(
            requester_admin=row.requester_admin,
            status="approved",
            asset_label=f"{row.asset_name or '-'} ({row.category or '-'})",
            acted_by=actor,
        )
    return _ok({"request": _serialize_return_request(row)}, "Return request approved")


@it_bp.route("/return-requests/<int:request_id>/reject", methods=["PATCH"])
@jwt_required()
def reject_return_request(request_id):
    actor = _current_admin()
    row = it_get(ITAssetReturnRequest, request_id)
    if not row:
        return _err("Return request not found", 404)
    if row.status != "pending":
        return _err("Only pending requests can be rejected")
    data = request.get_json(silent=True) or {}
    rejection_reason = (data.get("rejection_reason") or "").strip()
    if not rejection_reason:
        return _err("rejection_reason is required")
    err = _record_transition_or_err(
        action_code="REJECT_RETURN",
        remark=rejection_reason,
        actor_admin_id=actor.id if actor else None,
        asset_unit_id=row.asset_unit_id,
        software_license_id=row.software_license_id,
        inventory_item_id=row.inventory_item_id,
        from_status="pending",
        to_status="rejected",
        related={"return_request_id": row.id, "request_code": row.request_code},
    )
    if err:
        return err
    row.status = "rejected"
    row.approved_by_admin_id = actor.id if actor else None
    row.approved_at = utc_now()
    row.rejection_reason = rejection_reason
    db.session.commit()
    if row.requester_admin:
        send_it_return_request_status_email(
            requester_admin=row.requester_admin,
            status="rejected",
            asset_label=f"{row.asset_name or '-'} ({row.category or '-'})",
            acted_by=actor,
            rejection_reason=rejection_reason,
        )
    return _ok({"request": _serialize_return_request(row)}, "Return request rejected")


@it_bp.route("/return-requests/<int:request_id>/complete", methods=["PATCH"])
@jwt_required()
def complete_return_request(request_id):
    actor = _current_admin()
    row = it_get(ITAssetReturnRequest, request_id)
    if not row:
        return _err("Return request not found", 404)
    if row.status != "approved":
        return _err("Only approved requests can be completed")

    from .itam.transition_service import extract_remark_fields

    data = request.get_json(silent=True) or {}
    fields = extract_remark_fields(data)
    if not fields.get("remark"):
        fields["remark"] = (
            f"Physical receipt confirmed for return {row.request_code}. "
            f"Original reason: {(row.reason or '').strip() or 'n/a'}"
        )
    if not fields.get("condition_grade"):
        fields["condition_grade"] = "B"

    dest = (row.return_destination or "available").strip().lower()
    to_removed = dest == "removed_from_it"
    to_status = "not-working" if to_removed else "available"
    err = _record_transition_or_err(
        action_code="CHECKIN",
        remark=fields.get("remark"),
        reason_code=fields.get("reason_code"),
        condition_grade=fields.get("condition_grade"),
        actor_admin_id=actor.id if actor else None,
        asset_unit_id=row.asset_unit_id,
        software_license_id=row.software_license_id,
        inventory_item_id=row.inventory_item_id,
        from_status="approved",
        to_status=to_status,
        related={"return_request_id": row.id, "request_code": row.request_code},
    )
    if err:
        return err

    return_photos = row.photos_json if isinstance(row.photos_json, list) else []

    if row.asset_unit_id:
        unit = it_get(ITAssetUnit, row.asset_unit_id)
        if unit and unit.assigned_to_admin_id == row.requester_admin_id:
            if to_removed:
                unit.status = "not-working"
                db.session.add(
                    ITRemovedAsset(
                        removed_code=_next_code("RIT", ITRemovedAsset, "removed_code"),
                        asset_unit_id=unit.id,
                        inventory_item_id=unit.inventory_item_id,
                        owner_admin_id=row.requester_admin_id,
                        removed_by_admin_id=actor.id if actor else None,
                        name=row.asset_name or unit.asset_name or unit.unit_code,
                        category=row.category or unit.category,
                        reason=row.reason,
                        photos_json=return_photos,
                    )
                )
            else:
                unit.status = "available"
            unit.assigned_to_admin_id = None
            unit.assigned_at = None
            unit.repair_date = None
            assignment = ITAssetAssignment(
                assignment_type="return",
                assigned_to_admin_id=row.requester_admin_id,
                assigned_by_admin_id=actor.id if actor else None,
                asset_unit_id=unit.id,
                notes=f"Return request completed ({row.request_code})",
                assigned_at=utc_now(),
                unassigned_at=utc_now(),
            )
            db.session.add(assignment)
            _recalc_inventory_counts(unit.inventory_item_id)

    if row.software_license_id:
        lic = it_get(ITSoftwareLicense, row.software_license_id)
        if lic and lic.assigned_to_admin_id == row.requester_admin_id:
            if to_removed:
                lic.status = "not-working"
                db.session.add(
                    ITRemovedAsset(
                        removed_code=_next_code("RIT", ITRemovedAsset, "removed_code"),
                        inventory_item_id=lic.inventory_item_id,
                        owner_admin_id=row.requester_admin_id,
                        removed_by_admin_id=actor.id if actor else None,
                        name=row.asset_name or lic.name or lic.license_code,
                        category="Software",
                        reason=row.reason,
                        photos_json=return_photos,
                    )
                )
            else:
                lic.status = "available"
            lic.assigned_to_admin_id = None
            lic.assigned_at = None
            assignment = ITAssetAssignment(
                assignment_type="return",
                assigned_to_admin_id=row.requester_admin_id,
                assigned_by_admin_id=actor.id if actor else None,
                software_license_id=lic.id,
                notes=f"Return request completed ({row.request_code})",
                unassigned_at=utc_now(),
            )
            db.session.add(assignment)
            _recalc_inventory_counts(lic.inventory_item_id)

    if row.inventory_item_id and int(row.quantity or 0) > 0:
        item = it_get(ITInventoryItem, row.inventory_item_id)
        if item:
            qty = int(row.quantity or 0)
            qa = ITInventoryQuantityAssignment.query.filter_by(
                inventory_item_id=item.id,
                assigned_to_admin_id=row.requester_admin_id,
            ).first()
            if qa:
                qa.quantity = max(0, int(qa.quantity or 0) - qty)
                if qa.quantity == 0:
                    db.session.delete(qa)
            item.assigned_quantity = max(0, int(item.assigned_quantity or 0) - qty)
            if to_removed:
                db.session.add(
                    ITRemovedAsset(
                        removed_code=_next_code("RIT", ITRemovedAsset, "removed_code"),
                        inventory_item_id=item.id,
                        owner_admin_id=row.requester_admin_id,
                        removed_by_admin_id=actor.id if actor else None,
                        name=row.asset_name or item.name,
                        category=row.category or item.category,
                        reason=row.reason,
                        photos_json=return_photos,
                    )
                )
            else:
                item.available_quantity = int(item.available_quantity or 0) + qty

    row.status = "completed"
    row.receipt_confirmed_by_admin_id = actor.id if actor else None
    row.receipt_confirmed_at = utc_now()
    db.session.commit()
    if row.requester_admin:
        send_it_return_request_status_email(
            requester_admin=row.requester_admin,
            status="completed",
            asset_label=f"{row.asset_name or '-'} ({row.category or '-'})",
            acted_by=actor,
        )
    return _ok({"request": _serialize_return_request(row)}, "Return request completed")


@it_bp.route("/software/licenses/<int:license_id>/renew", methods=["PATCH"])
@jwt_required()
def renew_software_license(license_id):
    lic = it_get(ITSoftwareLicense, license_id)
    if not lic:
        return _err("Software license not found", 404)

    data = request.get_json(silent=True) or {}
    end_date = _parse_date(data.get("subscription_end") or data.get("new_expiry"))
    if not end_date:
        return _err("subscription_end/new_expiry is required in YYYY-MM-DD format")

    if lic.subscription_start and end_date <= lic.subscription_start:
        return _err("subscription_end must be after subscription_start")

    lic.subscription_end = end_date
    db.session.commit()
    return _ok({"license": _serialize_license(lic)}, "Software license renewed")


@it_bp.route("/assignments/inventory-quantity", methods=["POST"])
@jwt_required()
def assign_inventory_quantity():
    data = request.get_json(silent=True) or {}
    inventory_item_id = data.get("inventory_item_id")
    quantity = int(data.get("quantity") or 0)
    action = (data.get("action") or "assign").strip().lower()
    assigned_to_admin_id = data.get("assigned_to_admin_id")
    assigned_to_emp_id = (data.get("assigned_to_emp_id") or "").strip()
    target_admin = None
    actor = _current_admin()

    if not inventory_item_id or quantity < 1:
        return _err("inventory_item_id and quantity>=1 are required")
    if action not in ("assign", "return"):
        return _err("action must be assign or return")

    item = it_get(ITInventoryItem, inventory_item_id)
    if not item:
        return _err("Inventory item not found", 404)

    if action == "assign":
        if not assigned_to_admin_id and not assigned_to_emp_id:
            return _err("assigned_to_admin_id/assigned_to_emp_id is required for assign", 400)
        target_admin = get_admin_in_tenant(assigned_to_admin_id) if assigned_to_admin_id else None
        if not target_admin and assigned_to_emp_id:
            target_admin = tenant_admin_query().filter(
                db.func.lower(db.func.coalesce(Admin.emp_id, "")) == assigned_to_emp_id.lower()
            ).first()
        if not target_admin:
            return _err("Target admin not found", 404)
        if int(item.available_quantity or 0) < quantity:
            return _err("Not enough available quantity", 400)
        item.available_quantity = int(item.available_quantity or 0) - quantity
        item.assigned_quantity = int(item.assigned_quantity or 0) + quantity
        row = ITInventoryQuantityAssignment.query.filter_by(
            inventory_item_id=item.id,
            assigned_to_admin_id=target_admin.id,
        ).first()
        if row:
            row.quantity = int(row.quantity or 0) + quantity
        else:
            db.session.add(
                ITInventoryQuantityAssignment(
                    inventory_item_id=item.id,
                    assigned_to_admin_id=target_admin.id,
                    quantity=quantity,
                )
            )
    else:
        if int(item.assigned_quantity or 0) < quantity:
            return _err("Return quantity exceeds assigned quantity", 400)
        item.assigned_quantity = int(item.assigned_quantity or 0) - quantity
        item.available_quantity = int(item.available_quantity or 0) + quantity
        if assigned_to_admin_id or assigned_to_emp_id:
            target_admin = get_admin_in_tenant(assigned_to_admin_id) if assigned_to_admin_id else None
            if not target_admin and assigned_to_emp_id:
                target_admin = tenant_admin_query().filter(
                    db.func.lower(db.func.coalesce(Admin.emp_id, "")) == assigned_to_emp_id.lower()
                ).first()
            if target_admin:
                row = ITInventoryQuantityAssignment.query.filter_by(
                    inventory_item_id=item.id,
                    assigned_to_admin_id=target_admin.id,
                ).first()
                if row:
                    row.quantity = max(0, int(row.quantity or 0) - quantity)
                    if row.quantity == 0:
                        db.session.delete(row)

    db.session.commit()
    if action == "assign" and target_admin:
        try:
            ok, msg = send_it_assignment_notification(
                target_admin=target_admin,
                actor_admin=actor,
                assignment_kind="inventory_quantity",
                inventory_item=item,
                quantity=quantity,
            )
            if ok:
                current_app.logger.info(
                    "[IT assign_inventory_quantity] notification sent | inventory_item_id=%s target_admin_id=%s qty=%s msg=%s",
                    item.id,
                    target_admin.id,
                    quantity,
                    msg,
                )
            else:
                current_app.logger.warning(
                    "[IT assign_inventory_quantity] notification failed | inventory_item_id=%s target_admin_id=%s qty=%s msg=%s",
                    item.id,
                    target_admin.id,
                    quantity,
                    msg,
                )
        except Exception as e:
            current_app.logger.warning(
                "[IT assign_inventory_quantity] notification exception | inventory_item_id=%s target_admin_id=%s qty=%s err=%s",
                item.id,
                target_admin.id,
                quantity,
                e,
            )
    return _ok({"item": _serialize_inventory_item(item)}, "Inventory quantity updated")


_DEPLOY_INVENTORY_CATEGORIES = frozenset(
    {"Fixed Assets", "Office Assets", "Transport Assets", "Infrastructure Assets"}
)


def _inventory_deploy_mode(item):
    """Return 'stock' (qty), 'unit' (vehicle/equipment), or None."""
    inv_cat = (item.inventory_category or "").strip()
    if inv_cat not in _DEPLOY_INVENTORY_CATEGORIES:
        return None
    item_cat = (item.category or "").strip().lower()
    if _canonical_inv_cat(inv_cat) == "Fixed Assets":
        return "stock" if _is_qty_managed_inventory(item.category, item.inventory_category) else None
    if inv_cat == "Infrastructure Assets":
        if item_cat == "stock":
            return "stock"
        if item_cat == "equipment":
            return "unit"
        return None
    if inv_cat == "Transport Assets" and item_cat == "vehicle":
        return "unit"
    return None


def _serialize_office_deployment(row):
    unit = row.asset_unit if row.asset_unit_id else None
    unit_label = ""
    if unit:
        reg = (unit.serial_number or "").strip()
        make_model = " ".join(
            p for p in [(unit.make or "").strip(), (unit.model or "").strip()] if p
        ).strip()
        unit_label = reg or make_model or unit.unit_code or f"Unit #{unit.id}"
    return {
        "id": row.id,
        "inventoryItemId": row.inventory_item_id,
        "inventoryCategory": row.inventory_category or "",
        "assetUnitId": row.asset_unit_id,
        "unitLabel": unit_label,
        "quantity": int(row.quantity or 0),
        "deploymentLocation": row.deployment_location or "",
        "custodianName": row.custodian_name or "",
        "custodianEmpId": row.custodian_emp_id or "",
        "notes": row.notes or "",
        "createdAt": _iso(row.created_at),
    }


def _list_stock_deployments_impl():
    inventory_item_id = request.args.get("inventory_item_id")
    inventory_category = (request.args.get("inventory_category") or "").strip()
    q = ITOfficeStockDeployment.query.join(ITInventoryItem).filter(
        it_tenant_clause(ITInventoryItem),
        ITOfficeStockDeployment.quantity > 0,
        ITInventoryItem.inventory_category.in_(list(_DEPLOY_INVENTORY_CATEGORIES)),
    )
    if inventory_category:
        if inventory_category not in _DEPLOY_INVENTORY_CATEGORIES:
            return _err("Invalid inventory_category", 400)
        q = q.filter(ITInventoryItem.inventory_category.in_(_inv_cat_aliases(inventory_category)))
    if inventory_item_id:
        try:
            iid = int(inventory_item_id)
        except (TypeError, ValueError):
            return _err("inventory_item_id must be an integer", 400)
        q = q.filter(ITOfficeStockDeployment.inventory_item_id == iid)
    params = list_params()
    q = apply_search(
        q,
        (
            ITOfficeStockDeployment.deployment_location,
            ITOfficeStockDeployment.custodian_name,
            ITOfficeStockDeployment.custodian_emp_id,
        ),
        params.search,
    )
    newest = (ITOfficeStockDeployment.created_at.desc(),)
    if params.page is not None:
        newest += (ITOfficeStockDeployment.id.desc(),)
    order, err = _sort_or_err(
        params,
        {"created_at": ITOfficeStockDeployment.created_at, "location": ITOfficeStockDeployment.deployment_location},
        ITOfficeStockDeployment,
        newest,
    )
    if err:
        return err
    unit_labels = selectinload(ITOfficeStockDeployment.asset_unit).options(
        load_only(
            ITAssetUnit.id,
            ITAssetUnit.serial_number,
            ITAssetUnit.make,
            ITAssetUnit.model,
            ITAssetUnit.unit_code,
        )
    )
    rows, pagination = list_rows(q, ITOfficeStockDeployment, order, params, (unit_labels,), sort_ids_first=False)
    return _list_ok("deployments", [_serialize_office_deployment(r) for r in rows], pagination)


def _stock_deploy_impl():
    data = request.get_json(silent=True) or {}
    try:
        inventory_item_id = int(data.get("inventory_item_id"))
    except (TypeError, ValueError):
        return _err("inventory_item_id is required", 400)
    quantity = int(data.get("quantity") or 0)
    location = (data.get("deployment_location") or data.get("location") or "").strip()
    custodian_name = (data.get("custodian_name") or "").strip() or None
    custodian_emp_id = (data.get("custodian_emp_id") or "").strip() or None
    notes = (data.get("notes") or "").strip() or None
    asset_unit_id = data.get("asset_unit_id")

    if not location:
        return _err("deployment_location is required", 400)

    item = it_get(ITInventoryItem, inventory_item_id)
    if not item:
        return _err("Inventory item not found", 404)

    mode = _inventory_deploy_mode(item)
    if not mode:
        return _err("This item cannot be issued from inventory", 400)

    actor = _current_admin()
    inv_cat = (item.inventory_category or "").strip()

    if mode == "unit":
        try:
            uid = int(asset_unit_id)
        except (TypeError, ValueError):
            return _err("asset_unit_id is required for vehicles and installed equipment", 400)
        unit = it_get(ITAssetUnit, uid)
        if not unit or unit.inventory_item_id != item.id:
            return _err("Asset unit not found for this item", 404)
        if (unit.status or "").lower() != "available":
            return _err("This unit is not available to issue", 400)
        quantity = 1
        err = _record_transition_or_err(
            action_code="DEPLOY",
            remark=notes or f"Deployed to {location}",
            actor_admin_id=actor.id if actor else None,
            asset_unit_id=unit.id,
            inventory_item_id=item.id,
            from_status=unit.status,
            to_status="Deployed",
            to_custody={
                "type": "LOCATION",
                "location": location,
                "custodian_name": custodian_name,
            },
        )
        if err:
            return err
        unit.status = "assigned"
        unit.assigned_to_admin_id = None
        unit.assigned_at = None
        err = _apply_lifecycle_or_err(
            unit,
            lifecycle_status="Deployed",
            custody_type="LOCATION",
            custody={
                "type": "LOCATION",
                "location": location,
                "custodian_name": custodian_name,
                "custodian_emp_id": custodian_emp_id,
            },
            actor_admin_id=actor.id if actor else None,
            dual_write_legacy=False,
        )
        if err:
            return err
        db.session.add(
            ITOfficeStockDeployment(
                inventory_item_id=item.id,
                inventory_category=inv_cat,
                asset_unit_id=unit.id,
                quantity=1,
                deployment_location=location,
                custodian_name=custodian_name,
                custodian_emp_id=custodian_emp_id,
                notes=notes,
                created_by_admin_id=actor.id if actor else None,
            )
        )
        _recalc_inventory_counts(item.id)
        db.session.commit()
        return _ok(
            {"item": _serialize_inventory_item(item), "unit": _serialize_asset_unit(unit)},
            f"Issued to {location}",
            201,
        )

    if quantity < 1:
        return _err("quantity must be >= 1", 400)
    if int(item.available_quantity or 0) < quantity:
        return _err("Not enough available quantity", 400)

    err = _record_transition_or_err(
        action_code="DEPLOY",
        remark=notes or f"Deployed qty {quantity} to {location}",
        actor_admin_id=actor.id if actor else None,
        inventory_item_id=item.id,
        from_status="InStock",
        to_status="Deployed",
        to_custody={
            "type": "LOCATION",
            "location": location,
            "quantity": quantity,
            "custodian_name": custodian_name,
        },
    )
    if err:
        return err

    item.available_quantity = int(item.available_quantity or 0) - quantity
    item.assigned_quantity = int(item.assigned_quantity or 0) + quantity
    db.session.add(
        ITOfficeStockDeployment(
            inventory_item_id=item.id,
            inventory_category=inv_cat,
            quantity=quantity,
            deployment_location=location,
            custodian_name=custodian_name,
            custodian_emp_id=custodian_emp_id,
            notes=notes,
            created_by_admin_id=actor.id if actor else None,
        )
    )
    db.session.commit()
    return _ok(
        {"item": _serialize_inventory_item(item)},
        f"Issued {quantity} to {location}",
        201,
    )


def _stock_return_impl():
    from .itam.transition_service import extract_remark_fields

    data = request.get_json(silent=True) or {}
    try:
        deployment_id = int(data.get("deployment_id"))
    except (TypeError, ValueError):
        return _err("deployment_id is required", 400)
    quantity = int(data.get("quantity") or 0)
    if quantity < 1:
        return _err("quantity must be >= 1", 400)

    row = it_get(ITOfficeStockDeployment, deployment_id)
    if not row:
        return _err("Deployment not found", 404)
    item = it_get(ITInventoryItem, row.inventory_item_id)
    if not item or _inventory_deploy_mode(item) is None:
        return _err("Inventory item not found", 404)
    if int(row.quantity or 0) < quantity:
        return _err("Return quantity exceeds issued quantity", 400)

    fields = extract_remark_fields(data)
    actor = _current_admin()
    remark = fields.get("remark") or (row.notes or "").strip() or f"Returned {quantity} from {row.deployment_location}"
    err = _record_transition_or_err(
        action_code="UNDEPLOY",
        remark=remark,
        reason_code=fields.get("reason_code"),
        condition_grade=fields.get("condition_grade"),
        actor_admin_id=actor.id if actor else None,
        asset_unit_id=row.asset_unit_id,
        inventory_item_id=item.id,
        from_status="Deployed",
        to_status="InStock",
        from_custody={
            "type": "LOCATION",
            "location": row.deployment_location,
            "deployment_id": row.id,
        },
        to_custody={"type": "NONE"},
        related={"quantity": quantity},
    )
    if err:
        return err

    if row.asset_unit_id:
        if quantity != 1:
            return _err("Return quantity must be 1 for a single vehicle or equipment unit", 400)
        unit = it_get(ITAssetUnit, row.asset_unit_id)
        if unit:
            unit.status = "available"
            unit.assigned_to_admin_id = None
            unit.assigned_at = None
            err = _apply_lifecycle_or_err(
                unit,
                lifecycle_status="InStock",
                custody_type="NONE",
                custody={"type": "NONE"},
                actor_admin_id=actor.id if actor else None,
                dual_write_legacy=False,
            )
            if err:
                return err
        db.session.delete(row)
        _recalc_inventory_counts(item.id)
    else:
        row.quantity = int(row.quantity or 0) - quantity
        if row.quantity <= 0:
            db.session.delete(row)
        item.assigned_quantity = max(0, int(item.assigned_quantity or 0) - quantity)
        item.available_quantity = int(item.available_quantity or 0) + quantity

    db.session.commit()
    return _ok(
        {"item": _serialize_inventory_item(item)},
        f"Returned {quantity} to available",
    )


@it_bp.route("/office-stock/deployments", methods=["GET"])
@it_bp.route("/inventory-stock/deployments", methods=["GET"])
@jwt_required()
def list_office_stock_deployments():
    return _list_stock_deployments_impl()


@it_bp.route("/office-stock/deploy", methods=["POST"])
@it_bp.route("/inventory-stock/deploy", methods=["POST"])
@jwt_required()
def office_stock_deploy():
    return _stock_deploy_impl()


@it_bp.route("/office-stock/return", methods=["POST"])
@it_bp.route("/inventory-stock/return", methods=["POST"])
@jwt_required()
def office_stock_return():
    return _stock_return_impl()


@it_bp.route("/tickets", methods=["GET"])
@jwt_required()
def list_tickets():
    status = (request.args.get("status") or "").strip()
    params = list_params()
    order, err = _sort_or_err(
        params, _TICKET_SORTS, ITSupportTicket, (ITSupportTicket.created_at.desc(), ITSupportTicket.id.desc())
    )
    if err:
        return err
    q = it_query(ITSupportTicket)
    if status:
        q = q.filter(ci_equals(ITSupportTicket.status, status))
    priority = (request.args.get("priority") or "").strip() if pagination_enabled() else ""
    if priority:
        q = q.filter(ci_equals(ITSupportTicket.priority, priority))
    q = apply_search(q, (ITSupportTicket.ticket_code, ITSupportTicket.title), params.search)
    rows, pagination = list_rows(
        q,
        ITSupportTicket,
        order,
        params,
        (selectinload(ITSupportTicket.requester_admin), selectinload(ITSupportTicket.assignee_admin)),
    )
    return _list_ok("tickets", [_serialize_ticket(t) for t in rows], pagination)


_TICKET_SORTS = {
    "created_at": ITSupportTicket.created_at,
    "updated_at": ITSupportTicket.updated_at,
    "priority": ITSupportTicket.priority,
    "status": ITSupportTicket.status,
}


@it_bp.route("/tickets", methods=["POST"])
@jwt_required()
def create_ticket():
    data = request.get_json(silent=True) or {}
    requester_admin_id = data.get("requester_admin_id")
    title = (data.get("title") or "").strip()
    description = (data.get("description") or "").strip()
    if not requester_admin_id or not title or not description:
        return _err("requester_admin_id, title and description are required")
    requester = get_admin_in_tenant(requester_admin_id)
    if not requester:
        return _err("Requester not found", 404)
    t = ITSupportTicket(
        tenant_id=get_current_tenant_id(),
        ticket_code=_next_code("TKT", ITSupportTicket, "ticket_code"),
        requester_admin_id=requester.id,
        assignee_admin_id=_existing_fk_id(Admin, data.get("assignee_admin_id")),
        title=title,
        description=description,
        priority=(data.get("priority") or "medium").strip().lower(),
        status=(data.get("status") or "pending").strip().lower(),
    )
    db.session.add(t)
    db.session.commit()
    return _ok({"ticket": _serialize_ticket(t)}, "Ticket created", 201)


@it_bp.route("/tickets/<int:ticket_id>/resolve", methods=["PATCH"])
@jwt_required()
def resolve_ticket(ticket_id):
    t = it_get(ITSupportTicket, ticket_id)
    if not t:
        return _err("Ticket not found", 404)
    t.status = "completed"
    t.resolved_at = utc_now()
    db.session.commit()
    return _ok({"ticket": _serialize_ticket(t)}, "Ticket resolved")


@it_bp.route("/removed-assets", methods=["GET"])
@jwt_required()
def list_removed_assets():
    params = list_params()
    order, err = _sort_or_err(
        params,
        {"removed_at": ITRemovedAsset.removed_at, "name": ITRemovedAsset.name},
        ITRemovedAsset,
        (ITRemovedAsset.removed_at.desc(), ITRemovedAsset.id.desc()),
    )
    if err:
        return err
    q = apply_search(
        it_query(ITRemovedAsset),
        (ITRemovedAsset.removed_code, ITRemovedAsset.name, ITRemovedAsset.category, ITRemovedAsset.reason),
        params.search,
    )
    options = (selectinload(ITRemovedAsset.owner_admin), selectinload(ITRemovedAsset.removed_by_admin))
    if params.summary:
        options += (defer(ITRemovedAsset.photos_json),)
    rows, pagination = list_rows(q, ITRemovedAsset, order, params, options)
    return _list_ok("removed_assets", [_serialize_removed_asset(r, params.summary) for r in rows], pagination)


@it_bp.route("/removed-assets", methods=["POST"])
@jwt_required()
def create_removed_asset():
    data = request.get_json(silent=True) or {}
    name = (data.get("name") or "").strip()
    if not name:
        return _err("name is required")
    photos, err = _stored_media(data.get("photos") or [], "removed")
    if err:
        return err
    r = ITRemovedAsset(
        removed_code=data.get("removed_code") or _next_code("RIT", ITRemovedAsset, "removed_code"),
        asset_unit_id=_existing_fk_id(ITAssetUnit, data.get("asset_unit_id")),
        inventory_item_id=_existing_fk_id(ITInventoryItem, data.get("inventory_item_id")),
        owner_admin_id=_existing_fk_id(Admin, data.get("owner_admin_id")),
        removed_by_admin_id=_existing_fk_id(Admin, data.get("removed_by_admin_id")),
        name=name,
        category=data.get("category"),
        reason=data.get("reason"),
        photos_json=photos,
        removed_at=_parse_dt(data.get("removed_at")) or utc_now(),
    )
    db.session.add(r)
    try:
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        from sqlalchemy.exc import IntegrityError

        if isinstance(e, IntegrityError):
            r.asset_unit_id = None
            r.inventory_item_id = None
            r.owner_admin_id = None
            r.removed_by_admin_id = None
            db.session.add(r)
            db.session.commit()
        else:
            raise
    return _ok({"removed_asset": _serialize_removed_asset(r)}, "Removed asset logged", 201)


@it_bp.route("/removed-assets/<int:removed_id>", methods=["DELETE"])
@jwt_required()
def delete_removed_asset(removed_id):
    row = it_get(ITRemovedAsset, removed_id)
    if not row:
        return _err("Removed asset not found", 404)
    db.session.delete(row)
    db.session.commit()
    return _ok(message="Removed asset record deleted")


@it_bp.route("/deleted-logs", methods=["GET"])
@jwt_required()
def list_deleted_logs():
    params = list_params()
    order, err = _sort_or_err(
        params,
        {"deleted_at": ITDeletedAssetLog.deleted_at, "asset_name": ITDeletedAssetLog.asset_name},
        ITDeletedAssetLog,
        (ITDeletedAssetLog.deleted_at.desc(), ITDeletedAssetLog.id.desc()),
    )
    if err:
        return err
    q = apply_search(
        it_query(ITDeletedAssetLog),
        (
            ITDeletedAssetLog.delete_code,
            ITDeletedAssetLog.asset_name,
            ITDeletedAssetLog.serial_number,
            ITDeletedAssetLog.reason,
        ),
        params.search,
    )
    rows, pagination = list_rows(
        q, ITDeletedAssetLog, order, params, (selectinload(ITDeletedAssetLog.deleted_by_admin),)
    )
    return _list_ok("logs", [_serialize_deleted_log(d) for d in rows], pagination)


@it_bp.route("/deleted-logs", methods=["POST"])
@jwt_required()
def create_deleted_log():
    data = request.get_json(silent=True) or {}
    asset_unit_id = _existing_fk_id(ITAssetUnit, data.get("asset_unit_id"))
    inventory_item_id = _existing_fk_id(ITInventoryItem, data.get("inventory_item_id"))
    deleted_by_admin_id = _existing_fk_id(Admin, data.get("deleted_by_admin_id"))

    d = ITDeletedAssetLog(
        delete_code=data.get("delete_code") or _next_code("DEL", ITDeletedAssetLog, "delete_code"),
        asset_unit_id=asset_unit_id,
        inventory_item_id=inventory_item_id,
        deleted_by_admin_id=deleted_by_admin_id,
        deleted_by_name=(data.get("deleted_by_name") or "").strip() or None,
        asset_name=data.get("asset_name"),
        category=data.get("category"),
        serial_number=data.get("serial_number"),
        reason=data.get("reason"),
        deleted_at=_parse_dt(data.get("deleted_at")) or utc_now(),
    )
    db.session.add(d)
    try:
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        from sqlalchemy.exc import IntegrityError

        # Stale unit/inventory ids from client cache — keep the history row without FKs.
        if isinstance(e, IntegrityError):
            d.asset_unit_id = None
            d.inventory_item_id = None
            d.deleted_by_admin_id = None
            db.session.add(d)
            db.session.commit()
        else:
            raise
    return _ok({"log": _serialize_deleted_log(d)}, "Deleted asset log created", 201)


@it_bp.route("/deleted-logs/<string:delete_code>", methods=["DELETE"])
@jwt_required()
def delete_deleted_log(delete_code):
    row = it_query(ITDeletedAssetLog).filter(
        db.func.lower(db.func.coalesce(ITDeletedAssetLog.delete_code, "")) == str(delete_code).lower()
    ).first()
    if not row:
        return _err("Deleted log not found", 404)
    db.session.delete(row)
    db.session.commit()
    return _ok(message="Deleted log removed")


@it_bp.route("/deleted-logs", methods=["DELETE"])
@jwt_required()
def delete_all_deleted_logs():
    it_query(ITDeletedAssetLog).delete(synchronize_session=False)
    db.session.commit()
    return _ok(message="All deleted logs removed")


@it_bp.route("/parcels/imports", methods=["GET"])
@jwt_required()
def list_parcel_imports():
    _ensure_parcel_name_columns_runtime()
    try:
        page = max(int(request.args.get("page", 1) or 1), 1)
    except (TypeError, ValueError):
        page = 1
    try:
        per_page = int(request.args.get("per_page", 50) or 50)
    except (TypeError, ValueError):
        per_page = 50
    per_page = min(max(per_page, 1), 100)

    pagination = _paginate_parcel_rows(ITParcelImport, page=page, per_page=per_page)
    return _ok(
        {
            "imports": [_serialize_parcel_import(r) for r in pagination["items"]],
            "page": pagination["page"],
            "per_page": pagination["per_page"],
            "total": pagination["total"],
            "pages": pagination["pages"],
            "has_next": pagination["has_next"],
            "has_prev": pagination["has_prev"],
        }
    )


@it_bp.route("/parcels/imports", methods=["POST"])
@jwt_required()
def create_parcel_import():
    _ensure_parcel_name_columns_runtime()
    data = request.get_json(silent=True) or {}
    source = (data.get("source") or data.get("from") or "").strip()
    asset_name = (data.get("asset_name") or data.get("assetName") or "").strip()
    if not source or not asset_name:
        return _err("source/from and asset_name/assetName are required")
    received_by_name = (
        (data.get("received_by_name") or data.get("receivedBy") or "").strip() or None
    )
    photos, err = _stored_media(data.get("photos") or [], "parcels")
    if err:
        return err
    row = ITParcelImport(
        import_code=data.get("import_code") or _next_code("IMP", ITParcelImport, "import_code"),
        source=source,
        asset_name=asset_name,
        count=int(data.get("count") or 1),
        id_no=data.get("id_no") or data.get("idNo"),
        received_by_admin_id=_existing_fk_id(Admin, data.get("received_by_admin_id")),
        received_by_name=received_by_name,
        received_at=_parse_dt(data.get("received_at") or data.get("date")) or utc_now(),
        photos_json=photos,
    )
    db.session.add(row)
    db.session.flush()
    actor = _current_admin()
    import_remark = (
        f"Parcel import {row.import_code} of {asset_name} "
        f"(qty {int(row.count or 1)}) from {source}"
    )
    err = _record_transition_or_err(
        action_code="RECEIVE",
        remark=import_remark,
        actor_admin_id=actor.id if actor else None,
        related={
            "kind": "parcel_import",
            "parcel_import_id": row.id,
            "import_code": row.import_code,
            "source": source,
            "asset_name": asset_name,
        },
    )
    if err:
        db.session.rollback()
        return err
    try:
        db.session.commit()
    except Exception as exc:
        db.session.rollback()
        return _err(f"Could not save parcel import: {exc}", 500)
    return _ok({"import": _serialize_parcel_import(row)}, "Parcel import created", 201)


@it_bp.route("/parcels/exports", methods=["GET"])
@jwt_required()
def list_parcel_exports():
    _ensure_parcel_name_columns_runtime()
    try:
        page = max(int(request.args.get("page", 1) or 1), 1)
    except (TypeError, ValueError):
        page = 1
    try:
        per_page = int(request.args.get("per_page", 50) or 50)
    except (TypeError, ValueError):
        per_page = 50
    per_page = min(max(per_page, 1), 100)

    pagination = _paginate_parcel_rows(ITParcelExport, page=page, per_page=per_page)
    return _ok(
        {
            "exports": [_serialize_parcel_export(r) for r in pagination["items"]],
            "page": pagination["page"],
            "per_page": pagination["per_page"],
            "total": pagination["total"],
            "pages": pagination["pages"],
            "has_next": pagination["has_next"],
            "has_prev": pagination["has_prev"],
        }
    )


_MONTH_SHORT = ("", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")


def _parse_summary_year():
    raw = request.args.get("year")
    if raw in (None, ""):
        return utc_now().year, None
    try:
        year = int(raw)
    except (TypeError, ValueError):
        return None, _err("year must be a number")
    if year < 2000 or year > 2100:
        return None, _err("year must be between 2000 and 2100")
    return year, None


def _month_int(value):
    try:
        month = int(value)
    except (TypeError, ValueError):
        return None
    if month < 1 or month > 12:
        return None
    return month


def _parcel_monthly_buckets(year):
    start = datetime(year, 1, 1)
    end = datetime(year + 1, 1, 1)
    months = [
        {
            "month": f"{year}-{str(i).zfill(2)}",
            "label": f"{_MONTH_SHORT[i]} {year}",
            "imports": {"shipments": 0, "quantity": 0},
            "exports": {"shipments": 0, "units": 0},
        }
        for i in range(1, 13)
    ]

    import_month = extract("month", ITParcelImport.received_at)
    import_rows = (
        it_query(ITParcelImport)
        .with_entities(
            import_month,
            func.count(ITParcelImport.id),
            func.coalesce(func.sum(ITParcelImport.count), 0),
        )
        .filter(ITParcelImport.received_at >= start, ITParcelImport.received_at < end)
        .group_by(import_month)
        .all()
    )
    for month_val, shipments, quantity in import_rows:
        month = _month_int(month_val)
        if month is None:
            continue
        months[month - 1]["imports"]["shipments"] = int(shipments or 0)
        months[month - 1]["imports"]["quantity"] = int(quantity or 0)

    export_month = extract("month", ITParcelExport.exported_at)
    export_rows = (
        it_query(ITParcelExport)
        .with_entities(export_month, func.count(ITParcelExport.id))
        .filter(ITParcelExport.exported_at >= start, ITParcelExport.exported_at < end)
        .group_by(export_month)
        .all()
    )
    for month_val, shipments in export_rows:
        month = _month_int(month_val)
        if month is None:
            continue
        months[month - 1]["exports"]["shipments"] = int(shipments or 0)

    unit_rows = (
        it_query(ITParcelExport)
        .join(ITParcelExportItem, ITParcelExportItem.parcel_export_id == ITParcelExport.id)
        .with_entities(export_month, func.count(ITParcelExportItem.id))
        .filter(ITParcelExport.exported_at >= start, ITParcelExport.exported_at < end)
        .group_by(export_month)
        .all()
    )
    for month_val, units in unit_rows:
        month = _month_int(month_val)
        if month is None:
            continue
        months[month - 1]["exports"]["units"] = int(units or 0)

    totals = {
        "imports": {
            "shipments": sum(m["imports"]["shipments"] for m in months),
            "quantity": sum(m["imports"]["quantity"] for m in months),
        },
        "exports": {
            "shipments": sum(m["exports"]["shipments"] for m in months),
            "units": sum(m["exports"]["units"] for m in months),
        },
    }
    return months, totals


@it_bp.route("/parcels/monthly-summary", methods=["GET"])
@jwt_required()
def parcel_monthly_summary():
    year, err = _parse_summary_year()
    if err:
        return err
    months, totals = _parcel_monthly_buckets(year)
    return _ok({"year": year, "months": months, "totals": totals})


@it_bp.route("/parcels/exports", methods=["POST"])
@jwt_required()
def create_parcel_export():
    _ensure_parcel_name_columns_runtime()
    data = request.get_json(silent=True) or {}
    destination = (data.get("destination") or data.get("to") or "").strip()
    if not destination:
        return _err("destination/to is required")
    items = data.get("items") or data.get("assets") or []
    if not isinstance(items, list) or not items:
        return _err("items/assets must be a non-empty array")

    exported_by_name = (
        (data.get("exported_by_name") or data.get("exportedBy") or "").strip() or None
    )
    inventory_category = (
        (data.get("inventory_category") or data.get("inventoryCategory") or "").strip() or None
    )
    photos, err = _stored_media(data.get("photos") or [], "parcels")
    if err:
        return err
    row = ITParcelExport(
        export_code=data.get("export_code") or _next_code("EXP", ITParcelExport, "export_code"),
        destination=destination,
        id_no=data.get("id_no") or data.get("idNo"),
        exported_by_admin_id=_existing_fk_id(Admin, data.get("exported_by_admin_id")),
        exported_by_name=exported_by_name,
        exported_at=_parse_dt(data.get("exported_at") or data.get("date")) or utc_now(),
        parcel_photos_json=photos,
        inventory_category=inventory_category,
    )
    db.session.add(row)
    db.session.flush()

    created_items = []
    inventory_export_qty = {}
    for it in items:
        asset_unit_id, inventory_item_id = _resolve_parcel_export_item_ids(it)
        make_val = (it.get("make") or "").strip() or None
        if not make_val and asset_unit_id:
            unit_for_make = it_get(ITAssetUnit, asset_unit_id)
            if unit_for_make and unit_for_make.make:
                make_val = (unit_for_make.make or "").strip() or None
        photo, photo_err = _stored_media(
            it.get("individual_photo") or it.get("individualPhoto"), "parcels"
        )
        if photo_err:
            db.session.rollback()
            return photo_err
        export_item = ITParcelExportItem(
            parcel_export_id=row.id,
            asset_unit_id=asset_unit_id,
            asset_name=(it.get("asset_name") or it.get("assetName") or "").strip(),
            serial_number=it.get("serial_number") or it.get("serialNo"),
            brand=it.get("brand"),
            make=make_val,
            model=it.get("model"),
            individual_photo_json=photo,
        )
        if not export_item.asset_name:
            continue
        db.session.add(export_item)
        linked_inventory_id = inventory_item_id
        if asset_unit_id:
            unit = it_get(ITAssetUnit, asset_unit_id)
            if unit:
                linked_inventory_id = unit.inventory_item_id
                unit.status = "exported"
                unit.exported_to = destination
                unit.exported_at = row.exported_at
                _recalc_inventory_counts(unit.inventory_item_id)
        elif inventory_item_id:
            inventory_export_qty[inventory_item_id] = (
                inventory_export_qty.get(inventory_item_id, 0) + 1
            )
        created_items.append((export_item, linked_inventory_id))

    if not created_items:
        db.session.rollback()
        return _err("No valid parcel items found")

    for inv_id, qty in inventory_export_qty.items():
        item = it_get(ITInventoryItem, inv_id)
        if not item or qty < 1:
            continue
        if _is_qty_managed_inventory(item.category, item.inventory_category):
            item.available_quantity = max(0, int(item.available_quantity or 0) - qty)
            item.total_quantity = max(0, int(item.total_quantity or 0) - qty)
        else:
            units = ITAssetUnit.query.filter_by(inventory_item_id=inv_id).count()
            if units == 0:
                item.available_quantity = max(0, int(item.available_quantity or 0) - qty)
                item.total_quantity = max(0, int(item.total_quantity or 0) - qty)

    actor = _current_admin()
    export_remark = f"Parcel export {row.export_code} to {destination}"
    if len(export_remark.strip()) < 20:
        export_remark = f"{export_remark} — shipment recorded for parcel log"
    for export_item, linked_inventory_id in created_items:
        err = _record_transition_or_err(
            action_code="EXPORT",
            remark=export_remark,
            reason_code="PARCEL_EXPORT",
            actor_admin_id=actor.id if actor else None,
            asset_unit_id=export_item.asset_unit_id,
            inventory_item_id=linked_inventory_id,
            from_status=None,
            to_status="exported",
            related={
                "kind": "parcel_export",
                "parcel_export_id": row.id,
                "export_code": row.export_code,
                "destination": destination,
                "asset_name": export_item.asset_name,
                "inventory_category": inventory_category,
            },
        )
        if err:
            db.session.rollback()
            return err

    try:
        db.session.commit()
    except Exception as exc:
        db.session.rollback()
        return _err(f"Could not save parcel export: {exc}", 500)
    row = ITParcelExport.query.get(row.id)
    return _ok({"export": _serialize_parcel_export(row)}, "Parcel export created", 201)


def _ensure_it_staff_admin():
    admin = current_admin()
    if not admin:
        return None, _err("Employee not found", 404)
    from flask_jwt_extended import get_jwt
    from .plan_features import is_it_role, jwt_is_module_staff

    claims = get_jwt() or {}
    if jwt_is_module_staff(is_it_role, claims):
        return admin, None
    return None, _err("IT or Inventory access required", 403)


def _noc_scope_for_department_admin(admin):
    from .plan_features import is_inventory_department
    return "inventory" if is_inventory_department(admin.emp_type or "") else "it"


@it_bp.route("/noc-requests", methods=["GET"])
@jwt_required()
def it_list_noc_department_requests():
    admin, err = _ensure_it_staff_admin()
    if err:
        return err
    status_raw = (request.args.get("status") or "All").strip()
    try:
        items = list_noc_requests("it", admin, status_raw)
        return _ok({"requests": items})
    except Exception:
        current_app.logger.exception("it_list_noc_department_requests")
        return (
            jsonify(
                {
                    "success": False,
                    "message": "Unable to load NOC requests. Check that database migrations include noc_department_requests.",
                }
            ),
            500,
        )


@it_bp.route("/noc-requests/<int:req_id>/upload", methods=["POST"])
@jwt_required()
def it_upload_noc_department_document(req_id):
    admin, err = _ensure_it_staff_admin()
    if err:
        return err
    file = request.files.get("file")
    out = upload_noc_document("it", admin, req_id, file)
    code = out.pop("http", 200)
    return jsonify({k: v for k, v in out.items()}), code


@it_bp.route("/noc-requests/<int:req_id>/download", methods=["GET"])
@jwt_required()
def it_download_noc_department_document(req_id):
    admin, err = _ensure_it_staff_admin()
    if err:
        return err
    out = download_noc_document("it", admin, req_id)
    if not out.get("success"):
        return jsonify({"success": False, "message": out.get("message", "Error")}), out.get("http", 400)
    from .pdf_watermark import send_download_file
    return send_download_file(
        path=out["path"],
        download_name=out["download_name"],
        as_attachment=True,
    )


@it_bp.route("/inventory/noc-requests", methods=["GET"])
@jwt_required()
def inventory_list_noc_department_requests():
    admin, err = _ensure_it_staff_admin()
    if err:
        return err
    status_raw = (request.args.get("status") or "All").strip()
    try:
        items = list_noc_requests("inventory", admin, status_raw)
        return _ok({"requests": items})
    except Exception:
        current_app.logger.exception("inventory_list_noc_department_requests")
        return (
            jsonify(
                {
                    "success": False,
                    "message": "Unable to load inventory NOC requests.",
                }
            ),
            500,
        )


@it_bp.route("/inventory/noc-requests/<int:req_id>/upload", methods=["POST"])
@jwt_required()
def inventory_upload_noc_department_document(req_id):
    admin, err = _ensure_it_staff_admin()
    if err:
        return err
    file = request.files.get("file")
    out = upload_noc_document("inventory", admin, req_id, file)
    code = out.pop("http", 200)
    return jsonify({k: v for k, v in out.items()}), code


@it_bp.route("/inventory/noc-requests/<int:req_id>/download", methods=["GET"])
@jwt_required()
def inventory_download_noc_department_document(req_id):
    admin, err = _ensure_it_staff_admin()
    if err:
        return err
    out = download_noc_document("inventory", admin, req_id)
    if not out.get("success"):
        return jsonify({"success": False, "message": out.get("message", "Error")}), out.get("http", 400)
    from .pdf_watermark import send_download_file
    return send_download_file(
        path=out["path"],
        download_name=out["download_name"],
        as_attachment=True,
    )


@it_bp.route("/official-email-requests", methods=["GET"])
@jwt_required()
def it_list_official_email_requests():
    """IT queue: HR requests to generate official emails."""
    admin, err = _ensure_it_staff_admin()
    if err:
        return err
    from .official_email_service import list_official_email_requests

    status = (request.args.get("status") or "all").strip()
    return _ok(list_official_email_requests(get_current_tenant_id(), status_filter=status))


@it_bp.route("/official-email-requests/<int:request_id>/respond", methods=["POST"])
@jwt_required()
def it_respond_official_email_request(request_id):
    """IT generates the official email and replies to HR."""
    admin, err = _ensure_it_staff_admin()
    if err:
        return err
    from .official_email_service import respond_official_email_request

    data = request.get_json(silent=True) or {}
    result, message, code = respond_official_email_request(
        request_id,
        admin,
        generated_email=data.get("generated_email") or "",
        it_remarks=data.get("it_remarks") or "",
    )
    if result is None:
        return _err(message, code)
    return _ok(result, message=result.get("message", "OK"), code=code)


@it_bp.route("/asset-assignment-requests", methods=["GET"])
@jwt_required()
def it_list_asset_assignment_requests():
    """IT queue: asset assignment requests for newly approved employees."""
    admin, err = _ensure_it_staff_admin()
    if err:
        return err
    from .asset_assignment_service import list_asset_assignment_requests

    status = (request.args.get("status") or "all").strip()
    return _ok(list_asset_assignment_requests(get_current_tenant_id(), status_filter=status))


@it_bp.route("/asset-assignment-requests/<int:request_id>/complete", methods=["POST"])
@jwt_required()
def it_complete_asset_assignment_request(request_id):
    """IT marks assets as assigned; HR is informed by email."""
    admin, err = _ensure_it_staff_admin()
    if err:
        return err
    from .asset_assignment_service import complete_asset_assignment_request

    data = request.get_json(silent=True) or {}
    result, message, code = complete_asset_assignment_request(
        request_id,
        admin,
        assigned_assets=data.get("assigned_assets") or "",
        it_remarks=data.get("it_remarks") or "",
    )
    if result is None:
        return _err(message, code)
    return _ok(result, message=result.get("message", "OK"), code=code)


@it_bp.route("/itam/meta", methods=["GET"])
@jwt_required()
def itam_meta():
    """
    P0/P1 read-only: ITAM contract version, feature flags (default OFF), remark policies,
    lifecycle mapping draft, and planned API surface.
    """
    from flask import current_app
    from .itam.contracts import api_contract_snapshot

    return jsonify({"success": True, **api_contract_snapshot(current_app.config)}), 200


@it_bp.route("/units/<int:unit_id>/transitions", methods=["POST"])
@jwt_required()
def create_unit_transition(unit_id):
    """
    P1 dedicated transitions API. When itam_transitions_v1 is OFF, returns 409 so
    clients use legacy endpoints. When ON, records transition; for status-changing
    actions also applies legacy status mutation (CHECKOUT/CHECKIN/repair paths still
    prefer existing assignment endpoints).
    """
    from flask import current_app
    from .itam.transition_service import (
        TransitionValidationError,
        extract_remark_fields,
        record_transition,
        serialize_transition,
        transitions_enabled,
    )

    if not transitions_enabled(current_app.config):
        return _err(
            "ITAM transitions are disabled. Set ITAM_TRANSITIONS_V1=1 or use legacy endpoints.",
            409,
        )

    unit = it_get(ITAssetUnit, unit_id)
    if not unit:
        return _err("Unit not found", 404)

    data = request.get_json(silent=True) or {}
    action_code = (data.get("action_code") or "").strip().upper()
    if not action_code:
        return _err("action_code is required")

    fields = extract_remark_fields(data)
    actor = _current_admin()
    to_status = (data.get("to_status") or "").strip() or None

    # NOTE / status-only actions: record without forcing assignment APIs
    try:
        row = record_transition(
            action_code=action_code,
            remark=fields.get("remark"),
            reason_code=fields.get("reason_code"),
            condition_grade=fields.get("condition_grade"),
            actor_admin_id=actor.id if actor else None,
            asset_unit_id=unit.id,
            inventory_item_id=unit.inventory_item_id,
            from_status=unit.status,
            to_status=to_status or unit.status,
            from_custody=data.get("from_custody"),
            to_custody=data.get("custody") or data.get("to_custody"),
            related=data.get("related"),
            attachments=data.get("attachments"),
            config=current_app.config,
            require=True,
        )
    except TransitionValidationError as exc:
        return _err(str(exc), 400)

    # Apply simple status flips for repair/quarantine/restore/note-adjacent
    status_map = {
        "SEND_REPAIR": "repair",
        "MARK_QUARANTINE": "not-working",
        "COMPLETE_REPAIR": "available",
        "LOST": "not-working",
    }
    if action_code in status_map and (to_status or status_map[action_code]):
        new_status = to_status or status_map[action_code]
        unit.status = new_status
        if new_status == "repair" and not unit.repair_date:
            unit.repair_date = utc_now()
        if new_status == "available":
            unit.repair_date = None
        _recalc_inventory_counts(unit.inventory_item_id)

    db.session.commit()
    return _ok(
        {"transition": serialize_transition(row), "unit": _serialize_asset_unit(unit)},
        "Transition recorded",
        201,
    )


@it_bp.route("/units/<int:unit_id>/timeline", methods=["GET"])
@jwt_required()
def unit_timeline(unit_id):
    """P2: paginated asset history. Requires itam_timeline_v1."""
    from flask import current_app
    from .itam.timeline_service import query_transitions, timeline_enabled

    if not timeline_enabled(current_app.config):
        return _err(
            "ITAM timeline is disabled. Set ITAM_TIMELINE_V1=1 to enable.",
            409,
        )

    unit = it_get(ITAssetUnit, unit_id)
    if not unit:
        return _err("Unit not found", 404)

    actions_raw = (request.args.get("action") or request.args.get("actions") or "").strip()
    actions = [a for a in actions_raw.split(",") if a.strip()] if actions_raw else None
    result = query_transitions(
        asset_unit_id=unit_id,
        actions=actions,
        q=request.args.get("q"),
        date_from=request.args.get("from"),
        date_to=request.args.get("to"),
        page=request.args.get("page", 1, type=int),
        limit=request.args.get("limit", 50, type=int),
    )
    return _ok(
        {
            **result,
            "unit": _serialize_asset_unit(unit),
        }
    )


@it_bp.route("/units/<int:unit_id>/timeline.csv", methods=["GET"])
@jwt_required()
def unit_timeline_csv(unit_id):
    """P2: export full timeline CSV for one unit."""
    from flask import Response, current_app
    from .itam.timeline_service import query_transitions, timeline_enabled, timeline_to_csv

    if not timeline_enabled(current_app.config):
        return _err(
            "ITAM timeline is disabled. Set ITAM_TIMELINE_V1=1 to enable.",
            409,
        )

    unit = it_get(ITAssetUnit, unit_id)
    if not unit:
        return _err("Unit not found", 404)

    # Export up to 5000 newest events
    result = query_transitions(asset_unit_id=unit_id, page=1, limit=5000)
    csv_text = timeline_to_csv(result.get("transitions") or [])
    filename = f"asset-{unit.unit_code or unit_id}-timeline.csv"
    return Response(
        csv_text,
        mimetype="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@it_bp.route("/units/<int:unit_id>/timeline.xlsx", methods=["GET"])
@jwt_required()
def unit_timeline_xlsx(unit_id):
    """Export full timeline Excel for one unit."""
    from flask import current_app
    from .itam.timeline_service import query_transitions, timeline_enabled, timeline_to_xlsx
    from .utility import send_excel_file

    if not timeline_enabled(current_app.config):
        return _err(
            "ITAM timeline is disabled. Set ITAM_TIMELINE_V1=1 to enable.",
            409,
        )

    unit = it_get(ITAssetUnit, unit_id)
    if not unit:
        return _err("Unit not found", 404)

    result = query_transitions(asset_unit_id=unit_id, page=1, limit=5000)
    buf = timeline_to_xlsx(result.get("transitions") or [])
    filename = f"asset-{unit.unit_code or unit_id}-timeline.xlsx"
    return send_excel_file(buf, filename)


@it_bp.route("/activity-log", methods=["GET"])
@jwt_required()
def it_activity_log():
    """Org-wide audit log. Always available (not gated on timeline flag)."""
    from .itam.activity_log_service import query_activity_log

    actions_raw = (request.args.get("action") or request.args.get("actions") or "").strip()
    actions = [a for a in actions_raw.split(",") if a.strip()] if actions_raw else None
    actor_raw = request.args.get("actor_id") or request.args.get("actor")
    actor_id = None
    if actor_raw:
        try:
            actor_id = int(actor_raw)
        except (TypeError, ValueError):
            actor_id = None
    unit_raw = request.args.get("unit_id") or request.args.get("asset_unit_id")
    unit_id = None
    if unit_raw:
        try:
            unit_id = int(unit_raw)
        except (TypeError, ValueError):
            unit_id = None
    item_raw = request.args.get("inventory_item_id") or request.args.get("item_id")
    item_id = None
    if item_raw:
        try:
            item_id = int(item_raw)
        except (TypeError, ValueError):
            item_id = None
    result = query_activity_log(
        scope=request.args.get("scope"),
        actions=actions,
        q=request.args.get("q"),
        date_from=request.args.get("from"),
        date_to=request.args.get("to"),
        inventory_category=request.args.get("inventory_category")
        or request.args.get("category"),
        actor_admin_id=actor_id,
        asset_unit_id=unit_id,
        inventory_item_id=item_id,
        page=request.args.get("page", 1, type=int),
        limit=request.args.get("limit", 50, type=int),
    )
    return _ok(result)


@it_bp.route("/activity-log.csv", methods=["GET"])
@jwt_required()
def it_activity_log_csv():
    from flask import Response
    from .itam.activity_log_service import activity_log_to_csv, query_activity_log

    actions_raw = (request.args.get("action") or request.args.get("actions") or "").strip()
    actions = [a for a in actions_raw.split(",") if a.strip()] if actions_raw else None
    result = query_activity_log(
        scope=request.args.get("scope"),
        actions=actions,
        q=request.args.get("q"),
        date_from=request.args.get("from"),
        date_to=request.args.get("to"),
        inventory_category=request.args.get("inventory_category")
        or request.args.get("category"),
        page=1,
        limit=5000,
    )
    csv_text = activity_log_to_csv(result)
    scope = (request.args.get("scope") or "all").strip() or "all"
    filename = f"it-activity-log-{scope}.csv"
    return Response(
        csv_text,
        mimetype="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@it_bp.route("/activity-log.xlsx", methods=["GET"])
@jwt_required()
def it_activity_log_xlsx():
    from .itam.activity_log_service import activity_log_to_xlsx, query_activity_log
    from .utility import send_excel_file

    actions_raw = (request.args.get("action") or request.args.get("actions") or "").strip()
    actions = [a for a in actions_raw.split(",") if a.strip()] if actions_raw else None
    result = query_activity_log(
        scope=request.args.get("scope"),
        actions=actions,
        q=request.args.get("q"),
        date_from=request.args.get("from"),
        date_to=request.args.get("to"),
        inventory_category=request.args.get("inventory_category")
        or request.args.get("category"),
        page=1,
        limit=5000,
    )
    buf = activity_log_to_xlsx(result)
    scope = (request.args.get("scope") or "all").strip() or "all"
    filename = f"it-activity-log-{scope}.xlsx"
    return send_excel_file(buf, filename)


@it_bp.route("/units/<int:unit_id>/reviews", methods=["GET"])
@jwt_required()
def list_unit_reviews(unit_id):
    from .itam.asset_review_service import list_asset_reviews

    unit = it_get(ITAssetUnit, unit_id)
    if not unit:
        return _err("Unit not found", 404)
    reviews = list_asset_reviews(asset_unit_id=unit.id)
    return _ok({"reviews": reviews, "count": len(reviews)})


@it_bp.route("/units/<int:unit_id>/reviews", methods=["POST"])
@jwt_required()
def create_unit_review(unit_id):
    from .itam.asset_review_service import AssetReviewError, create_asset_review, serialize_review

    unit = it_get(ITAssetUnit, unit_id)
    if not unit:
        return _err("Unit not found", 404)
    actor = _current_admin()
    data = request.get_json(silent=True) or {}
    try:
        row = create_asset_review(
            asset_unit_id=unit.id,
            inventory_item_id=unit.inventory_item_id,
            actor_admin_id=actor.id if actor else None,
            data=data,
        )
        db.session.commit()
    except AssetReviewError as exc:
        db.session.rollback()
        return _err(str(exc), 400)
    except Exception:
        db.session.rollback()
        current_app.logger.exception("create_unit_review")
        return _err("Could not save review", 500)
    return _ok({"review": serialize_review(row)}, "Review saved", 201)


@it_bp.route("/inventory/items/<int:item_id>/reviews", methods=["GET"])
@jwt_required()
def list_inventory_reviews(item_id):
    from .itam.asset_review_service import list_asset_reviews

    item = it_get(ITInventoryItem, item_id)
    if not item:
        return _err("Inventory item not found", 404)
    reviews = list_asset_reviews(inventory_item_id=item.id)
    return _ok({"reviews": reviews, "count": len(reviews)})


@it_bp.route("/inventory/items/<int:item_id>/reviews", methods=["POST"])
@jwt_required()
def create_inventory_review(item_id):
    from .itam.asset_review_service import AssetReviewError, create_asset_review, serialize_review

    item = it_get(ITInventoryItem, item_id)
    if not item:
        return _err("Inventory item not found", 404)
    actor = _current_admin()
    data = request.get_json(silent=True) or {}
    try:
        row = create_asset_review(
            inventory_item_id=item.id,
            actor_admin_id=actor.id if actor else None,
            data=data,
        )
        db.session.commit()
    except AssetReviewError as exc:
        db.session.rollback()
        return _err(str(exc), 400)
    except Exception:
        db.session.rollback()
        current_app.logger.exception("create_inventory_review")
        return _err("Could not save review", 500)
    return _ok({"review": serialize_review(row)}, "Review saved", 201)


@it_bp.route("/software/licenses/<int:license_id>/timeline", methods=["GET"])
@jwt_required()
def software_timeline(license_id):
    from flask import current_app
    from .itam.timeline_service import query_transitions, timeline_enabled

    if not timeline_enabled(current_app.config):
        return _err(
            "ITAM timeline is disabled. Set ITAM_TIMELINE_V1=1 to enable.",
            409,
        )
    lic = it_get(ITSoftwareLicense, license_id)
    if not lic:
        return _err("Software license not found", 404)
    actions_raw = (request.args.get("action") or request.args.get("actions") or "").strip()
    actions = [a for a in actions_raw.split(",") if a.strip()] if actions_raw else None
    result = query_transitions(
        software_license_id=license_id,
        actions=actions,
        q=request.args.get("q"),
        date_from=request.args.get("from"),
        date_to=request.args.get("to"),
        page=request.args.get("page", 1, type=int),
        limit=request.args.get("limit", 50, type=int),
    )
    return _ok({**result, "license": _serialize_license(lic)})


@it_bp.route("/itam/backfill-assignment-history", methods=["POST"])
@jwt_required()
def itam_backfill_assignment_history():
    """
    P2 optional: seed transitions from ITAssetAssignment rows.
    Requires itam_timeline_v1 (or itam_transitions_v1) so ops can hydrate history.
    """
    from flask import current_app
    from .itam.flags import is_itam_flag_enabled
    from .itam.timeline_service import backfill_from_assignments

    cfg = current_app.config
    if not (
        is_itam_flag_enabled("itam_timeline_v1", cfg)
        or is_itam_flag_enabled("itam_transitions_v1", cfg)
    ):
        return _err(
            "Enable ITAM_TIMELINE_V1 or ITAM_TRANSITIONS_V1 before backfill.",
            409,
        )

    data = request.get_json(silent=True) or {}
    limit = int(data.get("limit") or request.args.get("limit") or 500)
    stats = backfill_from_assignments(limit=limit, config=cfg)
    return _ok(stats, "Assignment history backfill completed")


@it_bp.route("/itam/backfill-lifecycle", methods=["POST"])
@jwt_required()
def itam_backfill_lifecycle():
    """P3: hydrate lifecycle_status + custody records from legacy unit state."""
    from flask import current_app
    from .itam.flags import is_itam_flag_enabled
    from .itam.lifecycle_service import backfill_unit_lifecycle

    cfg = current_app.config
    if not is_itam_flag_enabled("itam_lifecycle_v1", cfg):
        return _err("Enable ITAM_LIFECYCLE_V1 before lifecycle backfill.", 409)

    data = request.get_json(silent=True) or {}
    limit = int(data.get("limit") or request.args.get("limit") or 500)
    stats = backfill_unit_lifecycle(limit=limit, config=cfg)
    return _ok(stats, "Lifecycle backfill completed")


def _opt_text(value, limit=None):
    text = str(value or "").strip()
    if not text:
        return None
    return text[:limit] if limit else text


def _opt_date(value):
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return datetime.strptime(text[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def _opt_amount(value):
    if value is None or value == "":
        return None
    try:
        return round(float(value), 2)
    except (TypeError, ValueError):
        return None


def _opt_int(value):
    if value is None or value == "":
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _repair_bill_row(row):
    admin = get_admin_in_tenant(row.created_by_admin_id) if row.created_by_admin_id else None
    date_val = row.date_of_repairing.isoformat() if row.date_of_repairing else None
    amount = float(row.bill_amount) if row.bill_amount is not None else None
    return {
        "id": row.id,
        "engineer_name": row.engineer_name or "",
        "laptop_code": row.laptop_code or "",
        "issue_type": row.issue_type or "",
        "date_of_repairing": date_val,
        "amount_status": row.amount_status or "",
        "working_status": row.working_status or "",
        "bill_amount": amount,
        "repair_count": row.repair_count,
        "remarks": row.remarks or "",
        "file_url": present_itam_media(row.file_key) if row.file_key else None,
        "file_name": row.file_name or "",
        "created_by": (admin.full_name if admin and getattr(admin, "full_name", None) else None)
        or (admin.first_name if admin else None)
        or "",
        "created_at": isoformat_api(row.created_at) if row.created_at else None,
    }


def _ensure_repair_bills_table():
    from sqlalchemy import inspect, text

    with db.engine.begin() as conn:
        conn.execute(text("SET SESSION lock_wait_timeout=5"))
        names = inspect(conn).get_table_names()
        if "it_repair_bills" in set(names):
            return
        ITRepairBill.__table__.create(bind=conn, checkfirst=True)


def _repair_bills_filtered(month=None, q=None):
    query = it_query(ITRepairBill)
    month = str(month or "").strip()
    if month and len(month) >= 7:
        try:
            year_n = int(month[:4])
            month_n = int(month[5:7])
            query = query.filter(
                db.or_(
                    db.and_(
                        extract("year", ITRepairBill.date_of_repairing) == year_n,
                        extract("month", ITRepairBill.date_of_repairing) == month_n,
                    ),
                    db.and_(
                        ITRepairBill.date_of_repairing.is_(None),
                        extract("year", ITRepairBill.created_at) == year_n,
                        extract("month", ITRepairBill.created_at) == month_n,
                    ),
                )
            )
        except ValueError:
            pass
    q = str(q or "").strip()
    if q:
        like = f"%{q}%"
        query = query.filter(
            db.or_(
                ITRepairBill.engineer_name.ilike(like),
                ITRepairBill.laptop_code.ilike(like),
                ITRepairBill.issue_type.ilike(like),
                ITRepairBill.amount_status.ilike(like),
                ITRepairBill.working_status.ilike(like),
                ITRepairBill.remarks.ilike(like),
            )
        )
    return query.order_by(ITRepairBill.created_at.desc(), ITRepairBill.id.desc())


def _prior_repair_count(laptop_code):
    code = _opt_text(laptop_code, 80)
    if not code:
        return 0
    return (
        it_query(ITRepairBill)
        .filter(func.lower(ITRepairBill.laptop_code) == code.lower())
        .count()
    )


@it_bp.route("/repair-bills", methods=["GET"])
@jwt_required()
def list_repair_bills():
    from sqlalchemy.exc import OperationalError, ProgrammingError

    try:
        _ensure_repair_bills_table()
    except (OperationalError, ProgrammingError, Exception):
        pass
    month = request.args.get("month")
    q = request.args.get("q")
    try:
        rows = _repair_bills_filtered(month=month, q=q).all()
    except (OperationalError, ProgrammingError):
        db.session.rollback()
        return _ok({
            "bills": [],
            "count": 0,
            "kpis": {
                "total_bills": 0,
                "total_amount": 0,
                "paid_amount": 0,
                "unpaid_count": 0,
            },
        })
    paid = 0.0
    unpaid_count = 0
    total_amount = 0.0
    for row in rows:
        amt = float(row.bill_amount) if row.bill_amount is not None else 0.0
        total_amount += amt
        status = (row.amount_status or "").strip().lower()
        if status == "paid":
            paid += amt
        elif status in ("unpaid", "partial", ""):
            unpaid_count += 1
    return _ok({
        "bills": [_repair_bill_row(r) for r in rows],
        "count": len(rows),
        "kpis": {
            "total_bills": len(rows),
            "total_amount": round(total_amount, 2),
            "paid_amount": round(paid, 2),
            "unpaid_count": unpaid_count,
        },
    })


@it_bp.route("/repair-bills/laptop-codes", methods=["GET"])
@jwt_required()
def repair_bill_laptop_codes():
    q = str(request.args.get("q") or "").strip()
    codes = set()
    bill_q = it_query(ITRepairBill).filter(ITRepairBill.laptop_code.isnot(None))
    if q:
        bill_q = bill_q.filter(ITRepairBill.laptop_code.ilike(f"%{q}%"))
    for (code,) in bill_q.with_entities(ITRepairBill.laptop_code).distinct().limit(40):
        if code:
            codes.add(code)
    unit_q = it_query(ITAssetUnit).filter(ITAssetUnit.unit_code.isnot(None))
    if q:
        unit_q = unit_q.filter(ITAssetUnit.unit_code.ilike(f"%{q}%"))
    for (code,) in unit_q.with_entities(ITAssetUnit.unit_code).distinct().limit(40):
        if code:
            codes.add(code)
    return _ok({"codes": sorted(codes)[:40]})


@it_bp.route("/repair-bills/prior-count", methods=["GET"])
@jwt_required()
def repair_bill_prior_count():
    code = request.args.get("laptop_code") or request.args.get("q")
    return _ok({"laptop_code": _opt_text(code, 80) or "", "prior_count": _prior_repair_count(code)})


@it_bp.route("/repair-bills", methods=["POST"])
@jwt_required()
def create_repair_bill():
    _ensure_repair_bills_table()
    data = request.get_json(silent=True) or {}
    laptop_code = _opt_text(data.get("laptop_code"), 80)
    repair_count = _opt_int(data.get("repair_count"))
    if repair_count is None and laptop_code:
        repair_count = _prior_repair_count(laptop_code) + 1
    file_raw = data.get("file") or data.get("attachment")
    file_name = _opt_text(data.get("file_name"), 220)
    file_key = None
    if file_raw:
        try:
            file_key = persist_itam_media(file_raw, "repair-bills")
        except ValueError as exc:
            return _err(str(exc), 400)
    row = ITRepairBill(
        engineer_name=_opt_text(data.get("engineer_name"), 160),
        laptop_code=laptop_code,
        issue_type=_opt_text(data.get("issue_type"), 120),
        date_of_repairing=_opt_date(data.get("date_of_repairing")),
        amount_status=_opt_text(data.get("amount_status"), 40),
        working_status=_opt_text(data.get("working_status"), 40),
        bill_amount=_opt_amount(data.get("bill_amount")),
        repair_count=repair_count,
        remarks=_opt_text(data.get("remarks")),
        file_key=file_key,
        file_name=file_name,
        created_by_admin_id=(current_admin().id if current_admin() else None),
    )
    db.session.add(row)
    db.session.commit()
    return _ok({"bill": _repair_bill_row(row)}, "Repair bill saved")


@it_bp.route("/repair-bills.xlsx", methods=["GET"])
@jwt_required()
def repair_bills_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from .utility import send_excel_file

    month = request.args.get("month")
    q = request.args.get("q")
    rows = _repair_bills_filtered(month=month, q=q).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Repair bills"
    headers = [
        "Engineer Name",
        "Laptop Code",
        "Issue Type",
        "Date of Repairing",
        "Amount Status",
        "Working Status",
        "Bill Amount",
        "Repair Count",
        "Remarks",
        "File",
        "Created By",
        "Created At",
    ]
    ws.append(headers)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="13877F")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        item = _repair_bill_row(row)
        ws.append([
            item["engineer_name"],
            item["laptop_code"],
            item["issue_type"],
            item["date_of_repairing"] or "",
            item["amount_status"],
            item["working_status"],
            item["bill_amount"] if item["bill_amount"] is not None else "",
            item["repair_count"] if item["repair_count"] is not None else "",
            item["remarks"],
            item["file_name"] or ("Yes" if item["file_url"] else ""),
            item["created_by"],
            (item["created_at"] or "")[:19],
        ])
    import io
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    label = (month or "all").replace("/", "-")
    return send_excel_file(buf, f"repair-bills-{label}.xlsx")


from . import it_inventory_screens  # noqa: E402,F401  (registers the inventory screen routes on it_bp)
from . import it_assets_screen  # noqa: E402,F401  (registers the Asset Management screen route on it_bp)
from .inventory_purchase import (  # noqa: E402
    cancel_request as cancel_inventory_purchase,
    create_request as create_inventory_purchase,
    finance_action as finance_inventory_purchase,
    list_requests as list_inventory_purchases,
    manager_action as manager_inventory_purchase,
    receive_request as receive_inventory_purchase,
)

it_bp.add_url_rule(
    "/inventory/purchase",
    endpoint="inventory_purchase_list",
    view_func=list_inventory_purchases,
    methods=["GET"],
    strict_slashes=False,
)
it_bp.add_url_rule(
    "/inventory/purchase",
    endpoint="inventory_purchase_create",
    view_func=create_inventory_purchase,
    methods=["POST"],
    strict_slashes=False,
)
it_bp.add_url_rule(
    "/inventory/purchase/<int:request_id>/manager-action",
    endpoint="inventory_purchase_manager",
    view_func=manager_inventory_purchase,
    methods=["POST"],
)
it_bp.add_url_rule(
    "/inventory/purchase/<int:request_id>/finance-action",
    endpoint="inventory_purchase_finance",
    view_func=finance_inventory_purchase,
    methods=["POST"],
)
it_bp.add_url_rule(
    "/inventory/purchase/<int:request_id>/receive",
    endpoint="inventory_purchase_receive",
    view_func=receive_inventory_purchase,
    methods=["POST"],
)
it_bp.add_url_rule(
    "/inventory/purchase/<int:request_id>/cancel",
    endpoint="inventory_purchase_cancel",
    view_func=cancel_inventory_purchase,
    methods=["POST"],
)
