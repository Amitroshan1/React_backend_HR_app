"""ITAM P2 — asset transition timeline query + CSV export (behind itam_timeline_v1)."""

from __future__ import annotations

import csv
import io
from datetime import datetime
from typing import Any, Optional

from .. import db
from ..models.Admin_models import Admin
from ..models.it_models import ITAssetAssignment, ITAssetTransition
from .actions import ACTION_LABELS, TransitionAction
from .flags import is_itam_flag_enabled
from .transition_service import record_transition, serialize_transition


def timeline_enabled(config=None) -> bool:
    return is_itam_flag_enabled("itam_timeline_v1", config)


def _parse_dt(value: Optional[str]) -> Optional[datetime]:
    if not value:
        return None
    raw = str(value).strip()
    try:
        if len(raw) == 10:
            return datetime.strptime(raw, "%Y-%m-%d")
        return datetime.fromisoformat(raw.replace("Z", "+00:00")).replace(tzinfo=None)
    except Exception:
        return None


def query_transitions(
    *,
    asset_unit_id: Optional[int] = None,
    software_license_id: Optional[int] = None,
    inventory_item_id: Optional[int] = None,
    actions: Optional[list[str]] = None,
    q: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    page: int = 1,
    limit: int = 50,
    inventory_category: Optional[str] = None,
    actor_admin_id: Optional[int] = None,
    parcel_only: bool = False,
    exclude_parcel: bool = False,
) -> dict[str, Any]:
    """Paginated timeline for a unit / license / catalog item, or org-wide activity log."""
    page = max(1, int(page or 1))
    limit = max(1, min(200, int(limit or 50)))

    query = ITAssetTransition.query
    if asset_unit_id is not None:
        query = query.filter(ITAssetTransition.asset_unit_id == int(asset_unit_id))
    if software_license_id is not None:
        query = query.filter(ITAssetTransition.software_license_id == int(software_license_id))
    if inventory_item_id is not None:
        query = query.filter(ITAssetTransition.inventory_item_id == int(inventory_item_id))
    if inventory_category:
        query = query.filter(
            ITAssetTransition.inventory_category == str(inventory_category).strip()
        )
    if actor_admin_id is not None:
        query = query.filter(ITAssetTransition.actor_admin_id == int(actor_admin_id))

    if actions:
        codes = [str(a).strip().upper() for a in actions if str(a).strip()]
        if codes:
            query = query.filter(ITAssetTransition.action_code.in_(codes))

    if parcel_only:
        # Parcel Log: exports + parcel imports only (exclude normal stock RECEIVE).
        query = query.filter(
            db.or_(
                ITAssetTransition.action_code == TransitionAction.EXPORT.value,
                ITAssetTransition.reason_code == "PARCEL_EXPORT",
                ITAssetTransition.remark.ilike("Parcel import%"),
                ITAssetTransition.remark.ilike("Parcel export%"),
            )
        )
    elif exclude_parcel:
        # Inventory / IT activity logs must not include parcel import/export events.
        query = query.filter(ITAssetTransition.action_code != TransitionAction.EXPORT.value)
        query = query.filter(
            db.or_(
                ITAssetTransition.reason_code.is_(None),
                ITAssetTransition.reason_code != "PARCEL_EXPORT",
            )
        )
        query = query.filter(
            db.or_(
                ITAssetTransition.remark.is_(None),
                db.and_(
                    ~ITAssetTransition.remark.ilike("Parcel import%"),
                    ~ITAssetTransition.remark.ilike("Parcel export%"),
                ),
            )
        )

    needle = (q or "").strip()
    if needle:
        like = f"%{needle}%"
        query = query.filter(
            db.or_(
                ITAssetTransition.remark.ilike(like),
                ITAssetTransition.reason_code.ilike(like),
                ITAssetTransition.action_code.ilike(like),
                ITAssetTransition.transition_code.ilike(like),
                ITAssetTransition.inventory_category.ilike(like),
                ITAssetTransition.from_status.ilike(like),
                ITAssetTransition.to_status.ilike(like),
            )
        )

    dt_from = _parse_dt(date_from)
    dt_to = _parse_dt(date_to)
    if dt_from:
        query = query.filter(ITAssetTransition.occurred_at >= dt_from)
    if dt_to:
        # inclusive end-of-day if date-only
        if date_to and len(str(date_to).strip()) == 10:
            dt_to = dt_to.replace(hour=23, minute=59, second=59)
        query = query.filter(ITAssetTransition.occurred_at <= dt_to)

    total = query.count()
    rows = (
        query.order_by(ITAssetTransition.occurred_at.desc(), ITAssetTransition.id.desc())
        .offset((page - 1) * limit)
        .limit(limit)
        .all()
    )

    actor_ids = {r.actor_admin_id for r in rows if r.actor_admin_id}
    actors = {}
    if actor_ids:
        for admin in Admin.query.filter(Admin.id.in_(list(actor_ids))).all():
            actors[admin.id] = {
                "id": admin.id,
                "name": (admin.first_name or admin.email or f"Admin #{admin.id}"),
                "empId": admin.emp_id,
                "email": admin.email,
            }

    transitions = []
    unit_ids = {r.asset_unit_id for r in rows if r.asset_unit_id}
    units_by_id = {}
    if unit_ids:
        from ..models.it_models import ITAssetUnit

        for unit in ITAssetUnit.query.filter(ITAssetUnit.id.in_(list(unit_ids))).all():
            units_by_id[unit.id] = unit

    for row in rows:
        item = serialize_transition(row)
        item["actionLabel"] = ACTION_LABELS.get(row.action_code, row.action_code)
        item["actor"] = actors.get(row.actor_admin_id)
        # Backfill device identity from live unit when older snapshots omitted fields.
        unit = units_by_id.get(row.asset_unit_id) if row.asset_unit_id else None
        if unit:
            item["assetName"] = item.get("assetName") or unit.asset_name
            item["serialNumber"] = item.get("serialNumber") or unit.serial_number
            item["unitCode"] = item.get("unitCode") or unit.unit_code
            item["brand"] = item.get("brand") or unit.brand
            item["make"] = item.get("make") or unit.make
            item["model"] = item.get("model") or unit.model
            item["laptopCode"] = item.get("laptopCode") or unit.model or unit.unit_code
            item["hwType"] = item.get("hwType") or unit.hw_type
            item["imei1"] = item.get("imei1") or unit.imei1
            item["imei2"] = item.get("imei2") or unit.imei2
            item["imei"] = item.get("imei") or unit.imei1 or unit.imei2
            item["projectCode"] = item.get("projectCode") or getattr(unit, "project_code", None)
            item["deviceLocation"] = item.get("deviceLocation") or getattr(
                unit, "device_location", None
            )
        transitions.append(item)

    latest = transitions[0] if transitions and page == 1 and not actions and not needle else None
    if latest is None and asset_unit_id and page == 1:
        latest_row = (
            ITAssetTransition.query.filter(ITAssetTransition.asset_unit_id == int(asset_unit_id))
            .order_by(ITAssetTransition.occurred_at.desc(), ITAssetTransition.id.desc())
            .first()
        )
        if latest_row:
            latest = serialize_transition(latest_row)
            latest["actionLabel"] = ACTION_LABELS.get(latest_row.action_code, latest_row.action_code)

    return {
        "transitions": transitions,
        "pagination": {
            "page": page,
            "limit": limit,
            "total": total,
            "totalPages": max(1, (total + limit - 1) // limit) if total else 1,
        },
        "latest": latest,
    }


def latest_by_unit_ids(unit_ids: list[int]) -> dict[int, dict]:
    """Map unit_id -> latest transition summary (one query)."""
    ids = [int(i) for i in unit_ids if i is not None]
    if not ids:
        return {}

    subq = (
        db.session.query(
            ITAssetTransition.asset_unit_id.label("uid"),
            db.func.max(ITAssetTransition.id).label("max_id"),
        )
        .filter(ITAssetTransition.asset_unit_id.in_(ids))
        .group_by(ITAssetTransition.asset_unit_id)
        .subquery()
    )
    rows = (
        db.session.query(ITAssetTransition)
        .join(subq, ITAssetTransition.id == subq.c.max_id)
        .all()
    )
    out = {}
    for row in rows:
        if not row.asset_unit_id:
            continue
        out[int(row.asset_unit_id)] = {
            "actionCode": row.action_code,
            "actionLabel": ACTION_LABELS.get(row.action_code, row.action_code),
            "remark": row.remark,
            "occurredAt": row.occurred_at.isoformat() if row.occurred_at else None,
            "transitionCode": row.transition_code,
        }
    return out


def timeline_to_csv(transitions: list[dict]) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(
        [
            "occurred_at",
            "transition_code",
            "action_code",
            "action_label",
            "asset_name",
            "hw_type",
            "brand",
            "model",
            "device_code",
            "imei",
            "serial_number",
            "device_location",
            "project_code",
            "inventory_category",
            "from_status",
            "to_status",
            "remark",
            "reason_code",
            "condition_grade",
            "actor_name",
            "actor_emp_id",
        ]
    )
    for t in transitions:
        actor = t.get("actor") or {}
        writer.writerow(
            [
                t.get("occurredAt") or "",
                t.get("transitionCode") or "",
                t.get("actionCode") or "",
                t.get("actionLabel") or "",
                t.get("assetName") or "",
                t.get("hwType") or "",
                t.get("brand") or "",
                t.get("model") or t.get("make") or "",
                t.get("laptopCode") or t.get("unitCode") or "",
                t.get("imei") or t.get("imei1") or t.get("imei2") or "",
                t.get("serialNumber") or "",
                t.get("deviceLocation") or "",
                t.get("projectCode") or "",
                t.get("inventoryCategory") or "",
                t.get("fromStatus") or "",
                t.get("toStatus") or "",
                t.get("remark") or "",
                t.get("reasonCode") or "",
                t.get("conditionGrade") or "",
                actor.get("name") or "",
                actor.get("empId") or "",
            ]
        )
    return buf.getvalue()


_XLSX_HEADERS = (
    "When",
    "Transition code",
    "Action code",
    "Action",
    "Asset",
    "Type",
    "Brand",
    "Model",
    "Device code",
    "IMEI",
    "Serial number",
    "Location",
    "Project",
    "Category",
    "From",
    "To",
    "Remark",
    "Reason",
    "Condition",
    "By",
    "Employee ID",
)


def timeline_to_xlsx(transitions: list[dict]) -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Activity log"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="3B82F6")
    ws.append(list(_XLSX_HEADERS))
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for t in transitions or []:
        actor = t.get("actor") or {}
        ws.append(
            [
                t.get("occurredAt") or "",
                t.get("transitionCode") or "",
                t.get("actionCode") or "",
                t.get("actionLabel") or "",
                t.get("assetName") or "",
                t.get("hwType") or "",
                t.get("brand") or "",
                t.get("model") or t.get("make") or "",
                t.get("laptopCode") or t.get("unitCode") or "",
                t.get("imei") or t.get("imei1") or t.get("imei2") or "",
                t.get("serialNumber") or "",
                t.get("deviceLocation") or "",
                t.get("projectCode") or "",
                t.get("inventoryCategory") or "",
                t.get("fromStatus") or "",
                t.get("toStatus") or "",
                t.get("remark") or "",
                t.get("reasonCode") or "",
                t.get("conditionGrade") or "",
                actor.get("name") or "",
                actor.get("empId") or "",
            ]
        )

    widths = [
        22, 16, 14, 22, 22, 12, 14, 14, 16, 18, 16, 16, 14, 14, 12, 12, 36, 12, 12, 18, 12,
    ]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def backfill_from_assignments(*, limit: int = 500, config=None) -> dict[str, int]:
    """
    Seed TransitionRecords from ITAssetAssignment history (best-effort).
    Skips assignments already linked via related_json.assignment_id.
    """
    limit = max(1, min(5000, int(limit or 500)))
    assignments = (
        ITAssetAssignment.query.filter(ITAssetAssignment.asset_unit_id.isnot(None))
        .order_by(ITAssetAssignment.id.asc())
        .limit(limit)
        .all()
    )

    # Preload existing backfill markers for scanned units
    unit_ids = {a.asset_unit_id for a in assignments if a.asset_unit_id}
    linked_assignment_ids = set()
    if unit_ids:
        prior_rows = (
            ITAssetTransition.query.filter(ITAssetTransition.asset_unit_id.in_(list(unit_ids)))
            .order_by(ITAssetTransition.id.desc())
            .limit(5000)
            .all()
        )
        for p in prior_rows:
            related = p.related_json if isinstance(p.related_json, dict) else {}
            aid = related.get("assignment_id")
            if aid is not None:
                linked_assignment_ids.add(int(aid))

    created = 0
    skipped = 0
    for asn in assignments:
        if asn.id in linked_assignment_ids:
            skipped += 1
            continue

        action = "CHECKOUT" if (asn.assignment_type or "").lower() in {"assign", ""} else "CHECKIN"
        remark = (asn.notes or "").strip() or (
            f"Backfilled from assignment #{asn.id} ({asn.assignment_type or 'assign'})"
        )
        # Satisfy remark min lengths
        if len(remark) < 10:
            remark = f"{remark} — historical assignment record #{asn.id}"

        try:
            cond = "B" if action == "CHECKIN" else None
            row = record_transition(
                action_code=action,
                remark=remark[:2000],
                reason_code=None,
                condition_grade=cond,
                actor_admin_id=asn.assigned_by_admin_id,
                asset_unit_id=asn.asset_unit_id,
                from_status="available" if action == "CHECKOUT" else "assigned",
                to_status="assigned" if action == "CHECKOUT" else "available",
                from_custody={"type": "NONE"}
                if action == "CHECKOUT"
                else {"type": "EMPLOYEE", "admin_id": asn.assigned_to_admin_id},
                to_custody={"type": "EMPLOYEE", "admin_id": asn.assigned_to_admin_id}
                if action == "CHECKOUT"
                else {"type": "NONE"},
                related={"assignment_id": asn.id, "backfill": True},
                config=config,
                require=True,
            )
            if row and asn.assigned_at:
                row.occurred_at = asn.assigned_at
            created += 1
            linked_assignment_ids.add(asn.id)
        except Exception:
            skipped += 1

    if created:
        db.session.commit()
    return {"created": created, "skipped": skipped, "scanned": len(assignments)}
